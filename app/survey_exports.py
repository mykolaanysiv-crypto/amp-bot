from __future__ import annotations

import json
from collections import Counter
from io import BytesIO
from textwrap import wrap
from typing import Any

from .models import Survey, SurveyQuestion, SurveyResponse, User


def parse_answers(response: SurveyResponse) -> dict[str, Any]:
    try:
        raw = json.loads(response.answers_json or "{}")
        return raw if isinstance(raw, dict) else {}
    except Exception:
        return {}


def build_survey_stats(
    questions: list[SurveyQuestion],
    responses: list[tuple[SurveyResponse, User]],
) -> tuple[dict[int, dict[str, Any]], dict[int, dict[str, Any]]]:
    stats: dict[int, dict[str, Any]] = {}
    decoded: dict[int, dict[str, Any]] = {}
    for q in questions:
        opts = [x.strip() for x in (q.options_text or "").splitlines() if x.strip()]
        stats[q.id] = {"options": opts, "counts": {x: 0 for x in opts}, "texts": []}
    for resp, _user in responses:
        answers = parse_answers(resp)
        decoded[resp.id] = answers
        for q in questions:
            value = answers.get(str(q.id))
            if q.question_type == "multiple":
                if isinstance(value, list):
                    for item in value:
                        if item in stats[q.id]["counts"]:
                            stats[q.id]["counts"][item] += 1
            elif q.question_type == "text":
                if value not in (None, ""):
                    stats[q.id]["texts"].append(str(value))
            elif value in stats[q.id]["counts"]:
                stats[q.id]["counts"][value] += 1
    return stats, decoded


def _answer_text(value: Any) -> str:
    if isinstance(value, list):
        return "; ".join(str(x) for x in value)
    if value is None:
        return ""
    return str(value)


def survey_excel(
    survey: Survey,
    questions: list[SurveyQuestion],
    responses: list[tuple[SurveyResponse, User]],
    stats: dict[int, dict[str, Any]],
    decoded: dict[int, dict[str, Any]],
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Зведення"
    ws.append(["АМПасадори — результати опитування", ""])
    ws.append(["Опитування", survey.title])
    ws.append(["Статус", survey.status])
    ws.append(["XP за проходження", int(survey.xp_reward or 0)])
    ws.append(["Кількість питань", len(questions)])
    ws.append(["Кількість респондентів", len(responses)])
    ws.append(["Дедлайн", survey.ends_at.strftime("%d.%m.%Y %H:%M") if survey.ends_at else "Без дедлайну"])
    ws.append(["Опис", survey.description or ""])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70

    qws = wb.create_sheet("Питання")
    qws.append(["№", "Питання", "Тип", "Обов'язкове", "Варіанти", "Фото"])
    type_labels = {"single": "Один варіант", "multiple": "Декілька варіантів", "text": "Текст"}
    for idx, q in enumerate(questions, 1):
        qws.append([
            idx,
            q.text,
            type_labels.get(q.question_type, q.question_type),
            "Так" if q.required else "Ні",
            " | ".join(x.strip() for x in (q.options_text or "").splitlines() if x.strip()),
            "Так" if getattr(q, "image_path", None) else "Ні",
        ])

    rws = wb.create_sheet("Респонденти")
    headers = ["ID", "Учасник", "Завершено", "XP"] + [f"{idx}. {q.text}" for idx, q in enumerate(questions, 1)]
    rws.append(headers)
    for resp, user in responses:
        answers = decoded.get(resp.id, {})
        rws.append([
            f"АМП-{user.id:04d}",
            user.full_name,
            resp.completed_at.strftime("%d.%m.%Y %H:%M"),
            int(resp.xp_awarded or 0),
            *[_answer_text(answers.get(str(q.id))) for q in questions],
        ])

    for idx, q in enumerate(questions, 1):
        title = f"П{idx}"[:31]
        sh = wb.create_sheet(title)
        sh.append([f"{idx}. {q.text}", ""])
        if q.question_type == "text":
            sh.append(["№", "Текстова відповідь"])
            for n, text in enumerate(stats[q.id]["texts"], 1):
                sh.append([n, text])
            sh.column_dimensions["A"].width = 10
            sh.column_dimensions["B"].width = 85
        else:
            sh.append(["Варіант", "Кількість"])
            for option, count in stats[q.id]["counts"].items():
                sh.append([option, count])
            if sh.max_row > 2:
                chart = BarChart()
                chart.title = f"Результат питання {idx}"
                chart.height = 8
                chart.width = 15
                chart.add_data(Reference(sh, min_col=2, min_row=2, max_row=sh.max_row), titles_from_data=True)
                chart.set_categories(Reference(sh, min_col=1, min_row=3, max_row=sh.max_row))
                sh.add_chart(chart, "D3")
            sh.column_dimensions["A"].width = 55
            sh.column_dimensions["B"].width = 16

    pale = "E8F8F9"
    teal = "06AEBB"
    thin = Side(style="thin", color="D7E5E8")
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=teal)
        if sh.max_row >= 2 and sh.title not in {"Зведення"}:
            for cell in sh[2]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor=pale)
        sh.freeze_panes = "A2"
        for col_idx in range(1, min(sh.max_column, 20) + 1):
            current = sh.column_dimensions[get_column_letter(col_idx)].width
            if not current or current < 12:
                sh.column_dimensions[get_column_letter(col_idx)].width = 18
    if rws.max_column > 4:
        for c in range(5, rws.max_column + 1):
            rws.column_dimensions[get_column_letter(c)].width = 38
    rws.column_dimensions["B"].width = 32

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _wrapped(text: str, width: int = 55) -> str:
    return "\n".join(wrap(str(text), width=width, break_long_words=False, break_on_hyphens=False))


def _draw_question(ax, q: SurveyQuestion, stat: dict[str, Any], response_count: int) -> None:
    if q.question_type == "text":
        ax.axis("off")
        texts = stat.get("texts", [])
        if not texts:
            ax.text(0.5, 0.5, "Текстових відповідей поки немає", ha="center", va="center", fontsize=13)
            return
        y = 0.96
        for idx, text in enumerate(texts[:12], 1):
            ax.text(0.02, y, f"{idx}. {_wrapped(text, 80)}", fontsize=9, va="top")
            y -= max(0.065, 0.035 * (1 + len(str(text)) // 80))
            if y < 0.06:
                break
        if len(texts) > 12:
            ax.text(0.02, 0.02, f"Ще відповідей: {len(texts)-12}", fontsize=8)
        return
    items = list(stat.get("counts", {}).items())
    labels = [x[0] for x in items]
    values = [x[1] for x in items]
    if not labels:
        ax.text(0.5, 0.5, "Даних поки немає", ha="center", va="center")
        ax.axis("off")
        return
    y = list(range(len(labels)))
    ax.barh(y, values)
    ax.set_yticks(y)
    ax.set_yticklabels([_wrapped(x, 34) for x in labels], fontsize=9)
    ax.invert_yaxis()
    ax.grid(axis="x", alpha=.2)
    ax.set_xlabel("Кількість відповідей")
    for yi, val in zip(y, values):
        pct = (val * 100 / response_count) if response_count else 0
        ax.text(val, yi, f" {val} ({pct:.0f}%)", va="center", fontsize=8)


def survey_question_png(
    survey: Survey,
    question: SurveyQuestion,
    stat: dict[str, Any],
    response_count: int,
) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams["font.family"] = "DejaVu Sans"
    bio = BytesIO()
    fig, ax = plt.subplots(figsize=(12, 7), dpi=150)
    fig.suptitle("АМПасадори • Результат опитування", fontsize=18, weight="bold", y=.98)
    fig.text(.06, .925, survey.title, fontsize=12, weight="bold")
    fig.text(.06, .885, _wrapped(question.text, 95), fontsize=10)
    _draw_question(ax, question, stat, response_count)
    fig.text(.06, .025, f"Респондентів: {response_count}", fontsize=8)
    fig.tight_layout(rect=[.04, .06, .98, .84])
    fig.savefig(bio, format="png", bbox_inches="tight")
    plt.close(fig)
    return bio.getvalue()


def survey_pdf(
    survey: Survey,
    questions: list[SurveyQuestion],
    responses: list[tuple[SurveyResponse, User]],
    stats: dict[int, dict[str, Any]],
    decoded: dict[int, dict[str, Any]],
) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    plt.rcParams["font.family"] = "DejaVu Sans"
    bio = BytesIO()
    with PdfPages(bio) as pdf:
        fig = plt.figure(figsize=(8.27, 11.69))
        fig.text(.08, .93, "АМПасадори — результати опитування", fontsize=19, weight="bold")
        fig.text(.08, .875, _wrapped(survey.title, 65), fontsize=15, weight="bold")
        fig.text(.08, .81, _wrapped(survey.description or "Без опису", 90), fontsize=9)
        rows = [
            ("Статус", survey.status),
            ("Питань", len(questions)),
            ("Респондентів", len(responses)),
            ("XP за проходження", int(survey.xp_reward or 0)),
            ("Дедлайн", survey.ends_at.strftime("%d.%m.%Y %H:%M") if survey.ends_at else "Без дедлайну"),
        ]
        y = .68
        for title, value in rows:
            fig.text(.10, y, title, fontsize=11)
            fig.text(.78, y, str(value), fontsize=11, weight="bold", ha="right")
            y -= .055
        fig.text(.08, .06, "АМПасадори • автоматично сформовано з даних опитування", fontsize=8)
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        for idx, q in enumerate(questions, 1):
            fig, ax = plt.subplots(figsize=(11.69, 8.27))
            fig.suptitle(f"Питання {idx}", fontsize=17, weight="bold", y=.96)
            fig.text(.07, .89, _wrapped(q.text, 105), fontsize=10.5)
            _draw_question(ax, q, stats[q.id], len(responses))
            fig.tight_layout(rect=[.04, .05, .98, .82])
            pdf.savefig(fig)
            plt.close(fig)

        if responses:
            # Compact respondent register; detailed answers are fully available in XLSX and web.
            fig, ax = plt.subplots(figsize=(11.69, 8.27))
            ax.axis("off")
            ax.set_title("Респонденти", fontsize=16, weight="bold", pad=20)
            table_rows = [
                [f"АМП-{user.id:04d}", user.full_name[:42], resp.completed_at.strftime("%d.%m.%Y %H:%M"), f"+{int(resp.xp_awarded or 0)}"]
                for resp, user in responses[:30]
            ]
            table = ax.table(
                cellText=table_rows,
                colLabels=["ID", "Учасник", "Завершено", "XP"],
                loc="upper center",
                cellLoc="left",
                colWidths=[.15, .48, .25, .12],
            )
            table.auto_set_font_size(False)
            table.set_fontsize(8.5)
            table.scale(1, 1.35)
            if len(responses) > 30:
                ax.text(.02, .05, f"У PDF показано перші 30 респондентів із {len(responses)}. Повний перелік є в Excel.", fontsize=8)
            fig.tight_layout()
            pdf.savefig(fig)
            plt.close(fig)
    return bio.getvalue()
