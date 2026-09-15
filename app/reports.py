from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from .models import (
    ActivityApplication, Event, EventRegistration, EventFeedback, Idea, OpportunityInterest, ParticipationStreak, Quest, RegistrationJourney,
    QuestParticipation, RequestCase, Season, StreakFreeze, Survey, SurveyResponse, User, UserBadge, VolunteerTask, VolunteerTaskParticipation,
    XPTransaction,
)
from .profile_data import CODE_TO_LABEL, gender_label, load_vulnerabilities
from .leagues import LEAGUES, league_for_xp
from .runtime_config import get_runtime_int
from .settlements import canonicalize_settlement_text, settlement_quality_report
from .time_utils import event_local_now

MONTHS_UA = ["січень","лютий","березень","квітень","травень","червень","липень","серпень","вересень","жовтень","листопад","грудень"]


def _next_month(year: int, month: int) -> tuple[int, int]:
    return (year + 1, 1) if month == 12 else (year, month + 1)


def resolve_report_period(period_type: str, *, year: int, month: int | None = None, start_month: int | None = None, end_month: int | None = None, quarter: int | None = None) -> tuple[datetime, datetime, str]:
    year = int(year)
    if not 2020 <= year <= 2100:
        raise ValueError("Некоректний рік")
    if period_type == "week":
        week = int(month or 1)
        if not 1 <= week <= 53: raise ValueError("Некоректний номер тижня")
        try:
            start_day = date.fromisocalendar(year, week, 1)
        except ValueError as exc:
            raise ValueError("Некоректний ISO-тиждень") from exc
        start_dt = datetime.combine(start_day, datetime.min.time())
        return start_dt, start_dt + timedelta(days=7), f"{week} тиждень {year}"
    if period_type == "month":
        m = int(month or 1)
        if not 1 <= m <= 12: raise ValueError("Некоректний місяць")
        ny, nm = _next_month(year, m)
        return datetime(year,m,1), datetime(ny,nm,1), f"{MONTHS_UA[m-1]} {year}"
    if period_type == "months":
        sm, em = int(start_month or 1), int(end_month or 12)
        if not (1 <= sm <= 12 and 1 <= em <= 12 and sm <= em): raise ValueError("Некоректний діапазон місяців")
        ny, nm = _next_month(year, em)
        label = f"{MONTHS_UA[sm-1]}–{MONTHS_UA[em-1]} {year}" if sm != em else f"{MONTHS_UA[sm-1]} {year}"
        return datetime(year,sm,1), datetime(ny,nm,1), label
    if period_type == "quarter":
        q = int(quarter or 1)
        if q not in {1,2,3,4}: raise ValueError("Некоректний квартал")
        sm = (q-1)*3+1; em=sm+2; ny,nm=_next_month(year,em)
        return datetime(year,sm,1), datetime(ny,nm,1), f"{q} квартал {year}"
    if period_type == "year":
        return datetime(year,1,1), datetime(year+1,1,1), f"{year} рік"
    raise ValueError("Невідомий тип періоду")


def _age(birth: date | None, on: date) -> int | None:
    if not birth: return None
    return on.year-birth.year-((on.month,on.day)<(birth.month,birth.day))


def _age_group(age: int | None) -> str:
    if age is None: return "Не зазначено"
    if age < 14: return "До 14"
    if age <=17: return "14–17"
    if age <=21: return "18–21"
    if age <=25: return "22–25"
    if age <=30: return "26–30"
    if age <=35: return "31–35"
    return "36+"


async def build_period_report(session: AsyncSession, start: datetime, end: datetime, label: str, *, reveal_sensitive_counts: bool = False) -> dict[str, Any]:
    generated_at = event_local_now()
    privacy_threshold = await get_runtime_int(session, "privacy.suppression_threshold")
    checkin_close_minutes = await get_runtime_int(session, "events.checkin_close_after_minutes")
    users=list((await session.scalars(select(User))).all())
    events=list((await session.scalars(select(Event).where(Event.starts_at>=start, Event.starts_at<end))).all())
    report_events=[e for e in events if e.status not in {"cancelled","draft"} and not e.cancelled_at]
    upcoming_events=[e for e in report_events if e.starts_at > generated_at]
    completed_events=[
        e for e in report_events
        if e.starts_at <= generated_at
        and (e.status == "completed" or e.starts_at + timedelta(minutes=checkin_close_minutes) < generated_at)
    ]
    in_progress_events=[e for e in report_events if e not in completed_events and e not in upcoming_events]
    completed_event_ids={e.id for e in completed_events}
    event_ids=[e.id for e in report_events]
    regs=list((await session.scalars(select(EventRegistration).where(EventRegistration.event_id.in_(event_ids) if event_ids else EventRegistration.id==-1))).all())
    attended=[r for r in regs if r.status=="attended" and r.event_id in completed_event_ids]
    quests=list((await session.scalars(select(Quest))).all())
    qparts=list((await session.scalars(select(QuestParticipation).where(QuestParticipation.status=="approved", QuestParticipation.approved_at>=start, QuestParticipation.approved_at<end))).all())
    tasks=list((await session.scalars(select(VolunteerTask))).all())
    tparts=list((await session.scalars(select(VolunteerTaskParticipation).where(VolunteerTaskParticipation.status=="approved", VolunteerTaskParticipation.approved_at>=start, VolunteerTaskParticipation.approved_at<end))).all())
    apps=list((await session.scalars(select(ActivityApplication).where(ActivityApplication.completed_at>=start, ActivityApplication.completed_at<end))).all())
    ideas=list((await session.scalars(select(Idea).where(Idea.created_at>=start, Idea.created_at<end))).all())
    implemented=list((await session.scalars(select(Idea).where(Idea.status=="implemented", Idea.updated_at>=start, Idea.updated_at<end))).all())
    requests=list((await session.scalars(select(RequestCase).where(RequestCase.created_at>=start, RequestCase.created_at<end))).all())
    resolved_requests=list((await session.scalars(select(RequestCase).where(RequestCase.resolved_at>=start, RequestCase.resolved_at<end))).all())
    txs=list((await session.scalars(select(XPTransaction).where(XPTransaction.created_at>=start, XPTransaction.created_at<end))).all())
    interests=list((await session.scalars(select(OpportunityInterest).where(OpportunityInterest.created_at>=start, OpportunityInterest.created_at<end, OpportunityInterest.status=="interested"))).all())
    all_surveys=list((await session.scalars(select(Survey))).all())
    surveys=[sv for sv in all_surveys if start <= (sv.starts_at or sv.created_at) < end and sv.status in {"published","closed"}]
    survey_responses=list((await session.scalars(select(SurveyResponse).where(SurveyResponse.completed_at>=start, SurveyResponse.completed_at<end))).all())
    user_badges=list((await session.scalars(select(UserBadge).where(UserBadge.awarded_at>=start, UserBadge.awarded_at<end))).all())
    feedback_rows=list((await session.scalars(select(EventFeedback).where(EventFeedback.completed_at>=start, EventFeedback.completed_at<end, EventFeedback.status=="completed"))).all())
    freeze_rows=list((await session.scalars(select(StreakFreeze).where(StreakFreeze.starts_at>=start, StreakFreeze.starts_at<end))).all())
    streak_rows=list((await session.scalars(select(ParticipationStreak))).all())
    active_season=await session.scalar(select(Season).where(Season.active==True).order_by(Season.starts_at.desc()))  # noqa: E712
    new_users=[u for u in users if start <= u.created_at < end]

    event_by={e.id:e for e in report_events}; task_by={t.id:t for t in tasks}
    event_attendance=Counter(r.event_id for r in attended)
    event_hours=sum(float(event_by[r.event_id].volunteer_hours or 0) for r in attended if r.event_id in event_by)
    task_hours=sum(float(task_by[p.task_id].hours_reward or 0) for p in tparts if p.task_id in task_by)
    activity_hours=sum(float(a.hours_reward or 0) for a in apps)

    engaged_ids=set(r.user_id for r in attended)
    engaged_ids.update(p.user_id for p in qparts)
    engaged_ids.update(p.user_id for p in tparts)
    engaged_ids.update(a.user_id for a in apps)
    engaged_ids.update(i.user_id for i in ideas)
    engaged_ids.update(c.user_id for c in requests)
    engaged_ids.update(r.user_id for r in survey_responses)
    cohort=[u for u in users if u.id in engaged_ids]

    gender=Counter(gender_label(u.gender) for u in cohort)
    age_reference=min(generated_at.date(), (end - timedelta(microseconds=1)).date())
    age=Counter(_age_group(_age(u.birth_date, age_reference)) for u in cohort)
    settlement=Counter(canonicalize_settlement_text(u.settlement) or "Не зазначено" for u in cohort)
    vuln_raw=Counter()
    for u in cohort:
        codes,_=load_vulnerabilities(u.vulnerability_categories)
        if not codes:
            vuln_raw["Не зазначено / не бажає повідомляти"]+=1
        else:
            for code in codes:
                vuln_raw["Інша категорія" if code=="other" else CODE_TO_LABEL.get(code,code)]+=1
    # Row-level sensitive data is never exported here. For ordinary staff, small
    # aggregate groups are suppressed; a superadmin may explicitly request exact
    # aggregate counts while still receiving no participant names/contacts.
    def _suppressed(count: int) -> bool:
        return (not reveal_sensitive_counts) and 0 < int(count) < privacy_threshold
    vuln=Counter({label:(0 if _suppressed(count) else int(count)) for label,count in vuln_raw.items()})
    vuln_display={label:(f"<{privacy_threshold}" if _suppressed(count) else str(int(count))) for label,count in vuln_raw.items()}

    event_rows=[]
    for e in sorted(report_events,key=lambda x:x.starts_at):
        if e in completed_events:
            timing_state, timing_label = "completed", "Завершена"
        elif e in upcoming_events:
            timing_state, timing_label = "upcoming", "Майбутня"
        else:
            timing_state, timing_label = "in_progress", "Триває"
        event_rows.append({
            "title":e.title,"date":e.starts_at.strftime("%d.%m.%Y %H:%M"),"location":e.location,"status":e.status,
            "timing_state":timing_state,"timing_label":timing_label,
            "attended":event_attendance.get(e.id,0) if timing_state=="completed" else 0,
            "registered":sum(1 for r in regs if r.event_id==e.id and r.status!="cancelled"),
        })
    future_attendance_anomalies=sum(1 for r in regs if r.status=="attended" and r.event_id not in completed_event_ids)

    # Monthly trend within selected interval.
    monthly=[]; y,m=start.year,start.month
    while datetime(y,m,1)<end:
        ny,nm=_next_month(y,m); ms=datetime(y,m,1); me=datetime(ny,nm,1)
        monthly.append({
            "label":f"{m:02d}.{y}",
            "events":sum(1 for e in completed_events if ms<=e.starts_at<me),
            "planned_events":sum(1 for e in report_events if ms<=e.starts_at<me),
            "visits":sum(1 for r in attended if (event_by.get(r.event_id) and ms<=event_by[r.event_id].starts_at<me)),
            "new_users":sum(1 for u in new_users if ms<=u.created_at<me),
            "xp":sum(max(0,int(t.amount or 0)) for t in txs if ms<=t.created_at<me),
            "survey_responses":sum(1 for r in survey_responses if ms<=r.completed_at<me),
            "badges":sum(1 for r in user_badges if ms<=r.awarded_at<me),
        })
        y,m=ny,nm

    # Reporting 2.0: choose a meaningful granularity instead of plotting a
    # one-point monthly line for short periods. Up to 31 days -> daily, up to
    # 120 days -> weekly, otherwise monthly.
    span_days = max(1, int((end - start).total_seconds() // 86400))
    if span_days <= 31:
        trend_granularity = "day"
        trend_granularity_label = "по днях"
        buckets=[]
        cursor=start
        while cursor < end:
            bucket_end=min(cursor+timedelta(days=1), end)
            buckets.append((cursor,bucket_end,cursor.strftime("%d.%m")))
            cursor=bucket_end
    elif span_days <= 120:
        trend_granularity = "week"
        trend_granularity_label = "по тижнях"
        buckets=[]
        cursor=start
        while cursor < end:
            bucket_end=min(cursor+timedelta(days=7), end)
            buckets.append((cursor,bucket_end,f"{cursor.strftime('%d.%m')}–{(bucket_end-timedelta(seconds=1)).strftime('%d.%m')}"))
            cursor=bucket_end
    else:
        trend_granularity = "month"
        trend_granularity_label = "по місяцях"
        buckets=[]
        y,m=start.year,start.month
        while datetime(y,m,1)<end:
            ny,nm=_next_month(y,m); bs=max(start,datetime(y,m,1)); be=min(end,datetime(ny,nm,1))
            buckets.append((bs,be,f"{m:02d}.{y}")); y,m=ny,nm

    trend=[]
    for bs,be,bucket_label in buckets:
        trend.append({
            "label":bucket_label,
            "events":sum(1 for e in completed_events if bs<=e.starts_at<be),
            "planned_events":sum(1 for e in report_events if bs<=e.starts_at<be),
            "visits":sum(1 for r in attended if (event_by.get(r.event_id) and bs<=event_by[r.event_id].starts_at<be)),
            "new_users":sum(1 for u in new_users if bs<=u.created_at<be),
            "xp":sum(max(0,int(t.amount or 0)) for t in txs if bs<=t.created_at<be),
            "survey_responses":sum(1 for r in survey_responses if bs<=r.completed_at<be),
            "badges":sum(1 for r in user_badges if bs<=r.awarded_at<be),
        })

    positive_xp=sum(int(t.amount or 0) for t in txs if int(t.amount or 0)>0)

    # Weekly badge dynamics inside the selected report period.
    badge_weekly_counter=Counter()
    for row in user_badges:
        monday=row.awarded_at.date()-timedelta(days=row.awarded_at.date().weekday())
        badge_weekly_counter[monday]+=1
    badge_weekly=[{"label":week.strftime("%d.%m.%Y"),"value":count} for week,count in sorted(badge_weekly_counter.items())]

    league_distribution={lg.title:0 for lg in LEAGUES}
    if active_season:
        xp_by_user=Counter()
        season_txs=list((await session.scalars(select(XPTransaction).where(XPTransaction.season_id==active_season.id))).all())
        for tx in season_txs: xp_by_user[tx.user_id]+=int(tx.amount or 0)
        for u in users:
            if getattr(u,"status",None)=="active":
                league_distribution[league_for_xp(xp_by_user.get(u.id,0)).title]+=1
    streak_snapshot={
        "weekly_active":sum(1 for r in streak_rows if int(r.weekly_streak or 0)>0),
        "super_active":sum(1 for r in streak_rows if int(r.event_streak or 0)>0),
        "recoverable":sum(1 for r in streak_rows if int(r.recoverable_event_streak or 0)>0),
        "best_weekly":max([int(r.weekly_best or 0) for r in streak_rows] or [0]),
        "best_event":max([int(r.event_best or 0) for r in streak_rows] or [0]),
    }
    def _pct_true(rows, attr: str) -> float:
        values=[getattr(row,attr,None) for row in rows if getattr(row,attr,None) is not None]
        return round(sum(1 for value in values if value)/len(values)*100,1) if values else 0.0

    ratings=[int(row.rating) for row in feedback_rows if row.rating is not None]
    feedback_avg_rating=round(sum(ratings)/len(ratings),2) if ratings else 0.0
    feedback_high_rating_pct=round(sum(1 for value in ratings if value>=4)/len(ratings)*100,1) if ratings else 0.0
    feedback_useful_pct=_pct_true(feedback_rows,"useful")
    feedback_new_knowledge_pct=_pct_true(feedback_rows,"new_knowledge")
    feedback_safe_pct=_pct_true(feedback_rows,"felt_safe")
    feedback_return_pct=_pct_true(feedback_rows,"would_return")
    outcomes={
        "responses":len(feedback_rows),
        "avg_rating":feedback_avg_rating,
        "high_rating_pct":feedback_high_rating_pct,
        "useful_pct":feedback_useful_pct,
        "new_knowledge_pct":feedback_new_knowledge_pct,
        "safe_pct":feedback_safe_pct,
        "return_pct":feedback_return_pct,
    }

    # v1.9.1 — donor/council-friendly advanced analytics.
    participant_users=[u for u in users if (getattr(u,"role",None) or "participant") in {"participant","ambassador"}]
    participant_ids={u.id for u in participant_users}
    all_event_regs=list((await session.scalars(select(EventRegistration).where(EventRegistration.status=="attended"))).all())
    all_event_ids={r.event_id for r in all_event_regs}
    all_events=list((await session.scalars(select(Event).where(Event.id.in_(all_event_ids) if all_event_ids else Event.id==-1))).all())
    all_event_by={e.id:e for e in all_events}
    visits_by_user=defaultdict(list)
    for reg in all_event_regs:
        if reg.user_id not in participant_ids:
            continue
        event=all_event_by.get(reg.event_id)
        stamp=event.starts_at if event and event.starts_at else (reg.confirmed_at or reg.checkin_at)
        if stamp:
            visits_by_user[reg.user_id].append(stamp)
    for values in visits_by_user.values(): values.sort()
    first_visit_ids={uid for uid,v in visits_by_user.items() if v}
    returned_ids={uid for uid,v in visits_by_user.items() if len(v)>=2}
    regular_ids={uid for uid,v in visits_by_user.items() if len(v)>=3 and len({(x.isocalendar().year,x.isocalendar().week) for x in v})>=3}
    ambassador_ids={u.id for u in participant_users if getattr(u,"role",None)=="ambassador"}
    cohort_funnel={
        "registered":len(participant_users),
        "first_visit":len(first_visit_ids),
        "returned":len(returned_ids),
        "regular":len(regular_ids),
        "ambassadors":len(regular_ids & ambassador_ids),
    }
    retention30=set(); retention90=set()
    for uid, visits in visits_by_user.items():
        if len(visits)<2: continue
        first=visits[0]
        if any(first < nxt <= first+timedelta(days=30) for nxt in visits[1:]): retention30.add(uid)
        if any(first < nxt <= first+timedelta(days=90) for nxt in visits[1:]): retention90.add(uid)
    retention_base=len(first_visit_ids)
    retention={
        "base":retention_base,
        "days30_count":len(retention30),
        "days30_pct":round(len(retention30)/retention_base*100,1) if retention_base else 0.0,
        "days90_count":len(retention90),
        "days90_pct":round(len(retention90)/retention_base*100,1) if retention_base else 0.0,
    }

    all_qparts=list((await session.scalars(select(QuestParticipation))).all())
    all_tparts=list((await session.scalars(select(VolunteerTaskParticipation))).all())
    all_apps=list((await session.scalars(select(ActivityApplication))).all())
    all_ideas=list((await session.scalars(select(Idea))).all())
    all_survey_responses=list((await session.scalars(select(SurveyResponse))).all())
    event_count=Counter(r.user_id for r in all_event_regs)
    quest_count=Counter(r.user_id for r in all_qparts if r.status=="approved")
    idea_count=Counter(r.user_id for r in all_ideas)
    survey_count=Counter(r.user_id for r in all_survey_responses)
    streak_by_user={row.user_id:int(row.weekly_streak or 0) for row in streak_rows}
    engagement_scores={}
    for u in participant_users:
        score=(min(streak_by_user.get(u.id,0)/4,1)*20 + min(event_count.get(u.id,0)/6,1)*25 + min(float(u.volunteer_hours or 0)/10,1)*20 + min(idea_count.get(u.id,0)/2,1)*10 + min(survey_count.get(u.id,0)/4,1)*10 + min(quest_count.get(u.id,0)/4,1)*15)
        engagement_scores[u.id]=int(round(score))
    engagement={
        "average":round(sum(engagement_scores.values())/len(engagement_scores),1) if engagement_scores else 0.0,
        "high":sum(1 for value in engagement_scores.values() if value>=75),
        "bands":{
            "0–24":sum(1 for value in engagement_scores.values() if 0<=value<=24),
            "25–49":sum(1 for value in engagement_scores.values() if 25<=value<=49),
            "50–74":sum(1 for value in engagement_scores.values() if 50<=value<=74),
            "75–100":sum(1 for value in engagement_scores.values() if 75<=value<=100),
        },
    }
    heatmap=[[0 for _ in range(24)] for _ in range(7)]
    heat_stamps=[]
    for reg in all_event_regs:
        event=all_event_by.get(reg.event_id); stamp=event.starts_at if event and event.starts_at else (reg.confirmed_at or reg.checkin_at)
        if stamp: heat_stamps.append(stamp)
    heat_stamps.extend((r.completed_at or r.approved_at) for r in all_qparts if r.status=="approved" and (r.completed_at or r.approved_at))
    heat_stamps.extend((r.submitted_at or r.approved_at) for r in all_tparts if r.status=="approved" and (r.submitted_at or r.approved_at))
    heat_stamps.extend((r.submitted_at or r.completed_at) for r in all_apps if r.status=="activity_completed" and (r.submitted_at or r.completed_at))
    heat_stamps.extend(r.completed_at for r in all_survey_responses if r.completed_at)
    heat_stamps.extend(i.created_at for i in all_ideas if i.created_at)
    for stamp in heat_stamps: heatmap[stamp.weekday()][stamp.hour]+=1

    # Reporting 2.0 conversion funnels. Registration funnel is cohort-based:
    # only journeys started in the selected period are included, and later
    # stages are counted only when they occurred before the report cut-off.
    cutoff=min(end, generated_at + timedelta(microseconds=1))
    journey_rows=list((await session.scalars(
        select(RegistrationJourney).where(RegistrationJourney.started_at>=start, RegistrationJourney.started_at<end)
    )).all())
    registration_funnel={
        "start":len(journey_rows),
        "consent":sum(1 for r in journey_rows if r.consent_at and r.consent_at < cutoff),
        "profile":sum(1 for r in journey_rows if r.profile_at and r.profile_at < cutoff),
        "submit":sum(1 for r in journey_rows if r.submitted_at and r.submitted_at < cutoff),
        "approved":sum(1 for r in journey_rows if r.approved_at and r.approved_at < cutoff),
        "first_activity":sum(1 for r in journey_rows if r.first_activity_at and r.first_activity_at < cutoff),
    }
    completed_regs=[r for r in regs if r.event_id in completed_event_ids]
    event_feedback_for_period=list((await session.scalars(
        select(EventFeedback).where(EventFeedback.event_id.in_(completed_event_ids) if completed_event_ids else EventFeedback.id==-1)
    )).all())
    event_xp_rows=list((await session.scalars(
        select(XPTransaction).where(
            XPTransaction.event_id.in_(completed_event_ids) if completed_event_ids else XPTransaction.id==-1,
            XPTransaction.category=="event", XPTransaction.amount>0,
        )
    )).all())
    event_conversion={
        "registered":sum(1 for r in completed_regs if r.status != "cancelled"),
        "checkin":sum(1 for r in completed_regs if r.checkin_at is not None or r.status in {"checked_in","attended"}),
        "attended":len(attended),
        "xp":len({r.user_id for r in event_xp_rows}),
        "feedback":sum(1 for r in event_feedback_for_period if r.status=="completed"),
    }

    settlement_quality=await settlement_quality_report(session)
    data_quality={
        "settlement_duplicate_groups":int(settlement_quality.get("duplicate_count",0)),
        "settlement_noncanonical":int(settlement_quality.get("noncanonical_count",0)),
        "profiles_without_settlement":sum(1 for u in users if not canonicalize_settlement_text(u.settlement)),
        "future_attendance_anomalies":future_attendance_anomalies,
    }
    data_quality["total_issues"] = sum(int(v or 0) for v in data_quality.values())

    indicator_definitions=[
        ("Завершені події","Події вибраного періоду, для яких завершилося операційне вікно або встановлено статус «Завершено»."),
        ("Унікальні залучені","Унікальні люди, які виконали хоча б одну підтверджену дію участі у вибраному періоді; сама реєстрація профілю не рахується."),
        ("Підтверджені відвідування","Лише підтверджена присутність на фактично завершених подіях; майбутні події не збільшують показник."),
        ("Потокові показники","Дії, що відбулися всередині вибраного періоду: відвідування, XP, квести, звернення, опитування тощо."),
        ("Моментні показники","Стан системи на момент формування звіту: активні/неактивні профілі, запити на відновлення та інші моментні значення."),
        ("Конверсія реєстрації","Когорта людей, які почали реєстрацію у вибраному періоді: старт → згода → профіль → надсилання → схвалення → перша активність."),
        ("Конверсія подій","Сукупний шлях реєстрацій на події періоду: заявка → відмітка → підтверджена участь → XP → завершений відгук."),
    ]

    summary={
        "events":len(completed_events),
        "events_planned":len(report_events),
        "events_upcoming":len(upcoming_events),
        "events_in_progress":len(in_progress_events),
        "unique_participants":len(engaged_ids),
        "visits":len(attended),
        "avg_attendance":round(len(attended)/len(completed_events),1) if completed_events else 0,
        "volunteer_hours":round(event_hours+task_hours+activity_hours,1),
        "tasks_completed":len(tparts),
        "quests_completed":len(qparts),
        "activities_completed":len(apps),
        "ideas_submitted":len(ideas),
        "ideas_implemented":len(implemented),
        "requests":len(requests),
        "requests_resolved":len(resolved_requests),
        "new_participants":len(new_users),
        "xp_awarded":positive_xp,
        "opportunity_interests":len(interests),
        "surveys_published":len(surveys),
        "survey_responses":len(survey_responses),
        "badges_awarded":len(user_badges),
        "streak_freeze_days":sum(int(r.days or 0) for r in freeze_rows),
        "active_profiles":sum(1 for u in users if getattr(u,"status",None)=="active"),
        "inactive_profiles":sum(1 for u in users if getattr(u,"status",None)=="inactive"),
        "deleted_profiles":sum(1 for u in users if getattr(u,"status",None)=="deleted"),
        "deleted_permanent_profiles":sum(1 for u in users if getattr(u,"status",None)=="deleted_permanent"),
        "restoration_requests_pending":sum(1 for u in users if getattr(u,"restoration_request_status",None)=="pending"),
        "probation_profiles":sum(1 for u in users if getattr(u,"probation_until",None) and getattr(u,"status",None)=="active"),
        "restored_profiles":sum(1 for u in users if getattr(u,"restored_at",None) is not None),
        "restoration_rejected":sum(1 for u in users if getattr(u,"restoration_request_status",None)=="rejected"),
        "participation_actions":len(attended)+len(qparts)+len(tparts)+len(apps)+len(survey_responses)+len(ideas),
        "avg_xp_per_engaged":round(positive_xp/len(engaged_ids),1) if engaged_ids else 0,
        "future_attendance_anomalies":future_attendance_anomalies,
        "feedback_responses":len(feedback_rows),
        "feedback_avg_rating":feedback_avg_rating,
        "feedback_high_rating_pct":feedback_high_rating_pct,
        "feedback_useful_pct":feedback_useful_pct,
        "feedback_new_knowledge_pct":feedback_new_knowledge_pct,
        "feedback_safe_pct":feedback_safe_pct,
        "feedback_return_pct":feedback_return_pct,
        "cohort_first_visit":cohort_funnel["first_visit"],
        "cohort_returned":cohort_funnel["returned"],
        "cohort_regular":cohort_funnel["regular"],
        "cohort_ambassadors":cohort_funnel["ambassadors"],
        "retention_30_pct":retention["days30_pct"],
        "retention_90_pct":retention["days90_pct"],
        "engagement_average":engagement["average"],
        "engagement_high":engagement["high"],
    }
    snapshot_keys={"active_profiles","inactive_profiles","deleted_profiles","deleted_permanent_profiles","restoration_requests_pending","probation_profiles","restored_profiles","restoration_rejected"}
    period_summary={k:v for k,v in summary.items() if k not in snapshot_keys}
    snapshot_summary={k:summary[k] for k in snapshot_keys}
    snapshot_summary["settlement_directory_profiles"] = sum(1 for u in users if canonicalize_settlement_text(u.settlement))
    return {"label":label,"start":start,"end":end,"generated_at":generated_at,"as_of":generated_at,"timezone":"Europe/Kyiv","summary":summary,"period_summary":period_summary,"snapshot_summary":snapshot_summary,"outcomes":outcomes,"events":event_rows,"monthly":monthly,"trend":trend,"trend_granularity":trend_granularity,"trend_granularity_label":trend_granularity_label,"registration_funnel":registration_funnel,"event_conversion":event_conversion,"data_quality":data_quality,"indicator_definitions":indicator_definitions,"age":age,"gender":gender,"settlement":settlement,"vulnerability":vuln,"vulnerability_display":vuln_display,"badge_weekly":badge_weekly,"league_distribution":league_distribution,"streak_snapshot":streak_snapshot,"cohort_funnel":cohort_funnel,"retention":retention,"engagement":engagement,"heatmap":heatmap,"heatmap_days":["Пн","Вт","Ср","Чт","Пт","Сб","Нд"],"privacy_note":(
        "Категорії вразливості подаються лише агреговано. Суперадміністратор бачить точні агреговані значення; ПІБ, контакти й списки конкретних осіб не формуються."
        if reveal_sensitive_counts else
        f"Категорії вразливості подаються лише агреговано. Значення 1–{privacy_threshold - 1} приховуються як <{privacy_threshold}; ПІБ, контакти й списки конкретних осіб не формуються."
    )}


def _style_excel(ws) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    teal="06AEBB"; deep="0B5B6C"; pale="EAF8FA"; alt="F7FBFC"; dark="173B43"; muted="6C858B"; white="FFFFFF"
    thin=Side(style="thin",color="D7E5E8")
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.tabColor=teal
    ws.page_margins.left=.35; ws.page_margins.right=.35; ws.page_margins.top=.55; ws.page_margins.bottom=.55
    ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.oddHeader.center.text="&BАМПасадори • Анисівський молодіжний простір"
    ws.oddHeader.center.size=10
    ws.oddHeader.center.font="Arial,Bold"
    ws.oddFooter.left.text="АМП • автоматичний звіт"
    ws.oddFooter.right.text="Сторінка &P з &N"
    for row in ws.iter_rows():
        for c in row:
            c.alignment=Alignment(vertical="top",wrap_text=True)
            c.border=Border(bottom=thin)
            c.font=Font(name="Arial",size=10,color=dark)
    # Brand the first row as a title/header row.
    for c in ws[1]:
        c.font=Font(name="Arial",bold=True,size=14,color=white)
        c.fill=PatternFill("solid",fgColor=deep)
        c.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[1].height=26
    # Detect secondary table headers used by the report sheets.
    header_names={"Показник","Подія","Місяць","Категорія","Ліга","Тиждень"}
    for row in ws.iter_rows(min_row=2):
        first=str(row[0].value or "")
        if first in header_names:
            for c in row:
                c.font=Font(name="Arial",bold=True,size=10,color=deep)
                c.fill=PatternFill("solid",fgColor=pale)
                c.alignment=Alignment(vertical="center",wrap_text=True)
        elif row[0].row % 2 == 0:
            for c in row:
                if c.value is not None and c.fill.fill_type is None:
                    c.fill=PatternFill("solid",fgColor=alt)
    ws.freeze_panes="A2"


def report_excel(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter
    wb=Workbook(); ws=wb.active; ws.title="Зведення"
    ws.append(["АМПасадори — автоматичний звіт",""])
    ws.append(["Період",data["label"]]); ws.append(["Дані станом на",f"{data['generated_at'].strftime('%d.%m.%Y %H:%M')} {data.get('timezone','Europe/Kyiv')}"]); ws.append(["Примітка",data["privacy_note"]]); ws.append([])
    labels={"events":"Завершені події","events_planned":"Усього заплановано подій","events_upcoming":"Майбутні події","events_in_progress":"Події у поточному вікні","unique_participants":"Унікальні залучені учасники","visits":"Підтверджені відвідування","avg_attendance":"Середня відвідуваність","volunteer_hours":"Волонтерські години","tasks_completed":"Виконані волонтерські задачі","quests_completed":"Підтверджені квести","activities_completed":"Підтверджені активності","ideas_submitted":"Подані ідеї","ideas_implemented":"Реалізовані ідеї","requests":"Звернення","requests_resolved":"Вирішені звернення","new_participants":"Нові учасники","xp_awarded":"Нараховано XP","opportunity_interests":"Позначки «Мені цікаво»","surveys_published":"Опубліковані опитування","survey_responses":"Проходження опитувань","badges_awarded":"Отримані бейджі","streak_freeze_days":"Днів заморозки серій","active_profiles":"Активні профілі","inactive_profiles":"Неактивні профілі","deleted_profiles":"Видалені профілі","deleted_permanent_profiles":"Видалені без відновлення","restoration_requests_pending":"Запити на відновлення","probation_profiles":"На 14-денному випробувальному строку","restored_profiles":"Відновлені профілі","restoration_rejected":"Відхилені запити на відновлення","participation_actions":"Дій участі","avg_xp_per_engaged":"Середній XP на залученого","future_attendance_anomalies":"Некоректні підтвердження участі поза завершеними подіями","settlement_directory_profiles":"Профілі з канонічним населеним пунктом","feedback_responses":"Зворотний зв’язок — відповідей","feedback_avg_rating":"Зворотний зв’язок — середня оцінка","feedback_high_rating_pct":"Висока оцінка 4–5, %","feedback_useful_pct":"Було корисно, %","feedback_new_knowledge_pct":"Нові знання, %","feedback_safe_pct":"Почувалися безпечно, %","feedback_return_pct":"Хочуть прийти ще, %","cohort_first_visit":"Когорта — прийшли 1 раз","cohort_returned":"Когорта — повернулися","cohort_regular":"Когорта — регулярні","cohort_ambassadors":"Когорта — АМПасадори","retention_30_pct":"Повернення за 30 днів, %","retention_90_pct":"Повернення за 90 днів, %","engagement_average":"Середній індекс залученості","engagement_high":"Залученість 75–100"}
    ws.append(["ПОТОКОВІ ПОКАЗНИКИ ЗА ПЕРІОД", ""]); ws.append(["Показник","Значення"])
    for k,v in data.get("period_summary", data["summary"]).items(): ws.append([labels.get(k,k),v])
    ws.append([]); ws.append(["МОМЕНТНІ ПОКАЗНИКИ СТАНОМ НА ДАТУ", f"{data['generated_at'].strftime('%d.%m.%Y %H:%M')} {data.get('timezone','Europe/Kyiv')}"]); ws.append(["Показник","Значення"])
    for k,v in data.get("snapshot_summary", {}).items(): ws.append([labels.get(k,k),v])
    ws.column_dimensions["A"].width=44; ws.column_dimensions["B"].width=24; _style_excel(ws)

    ew=wb.create_sheet("Події"); ew.append(["Подія","Дата/час","Локація","Стан у звіті","Системний статус","Зареєстровано","Відвідали"])
    for r in data["events"]: ew.append([r["title"],r["date"],r["location"],r.get("timing_label",r["status"]),r["status"],r["registered"],r["attended"]])
    for i,w in enumerate([42,20,26,18,16,16,14],1): ew.column_dimensions[get_column_letter(i)].width=w
    _style_excel(ew)

    mw=wb.create_sheet("Динаміка"); mw.append([f"Період ({data.get('trend_granularity_label','')})","Завершені події","Заплановані події","Відвідування","Нові учасники","XP","Опитування — відповіді","Отримані бейджі"])
    trend_rows=data.get("trend",data.get("monthly",[]))
    for r in trend_rows: mw.append([r["label"],r["events"],r.get("planned_events",r["events"]),r["visits"],r["new_users"],r["xp"],r.get("survey_responses",0),r.get("badges",0)])
    if len(trend_rows)>1:
        chart=LineChart(); chart.title=f"Динаміка відвідувань {data.get('trend_granularity_label','')}"; chart.add_data(Reference(mw,min_col=4,min_row=1,max_row=mw.max_row),titles_from_data=True); chart.set_categories(Reference(mw,min_col=1,min_row=2,max_row=mw.max_row)); mw.add_chart(chart,"J2")
    elif len(trend_rows)==1:
        mw["J2"]="Одна точка даних — лінійний графік не будується."
        mw["J3"]="Відвідування"; mw["K3"]=trend_rows[0]["visits"]
        mw["J4"]="Нові учасники"; mw["K4"]=trend_rows[0]["new_users"]
        mw["J5"]="XP"; mw["K5"]=trend_rows[0]["xp"]
    _style_excel(mw)

    fw=wb.create_sheet("Воронки")
    fw.append(["Конверсія реєстрації", "Кількість"]); fw.append(["Етап","Кількість"])
    reg_labels={"start":"Старт","consent":"Згода","profile":"Профіль","submit":"Надсилання","approved":"Схвалення","first_activity":"Перша активність"}
    for key in ["start","consent","profile","submit","approved","first_activity"]: fw.append([reg_labels[key],data.get("registration_funnel",{}).get(key,0)])
    fw.append([]); fw.append(["Конверсія участі у подіях", "Кількість"]); fw.append(["Етап","Кількість"])
    event_labels={"registered":"Заявки / реєстрації","checkin":"Відмітка","attended":"Підтверджена участь","xp":"XP нараховано","feedback":"Завершений відгук"}
    for key in ["registered","checkin","attended","xp","feedback"]: fw.append([event_labels[key],data.get("event_conversion",{}).get(key,0)])
    fw.column_dimensions["A"].width=34; fw.column_dimensions["B"].width=18; _style_excel(fw)

    dqw=wb.create_sheet("Якість даних")
    dqw.append(["Якість даних", "Кількість"]); dqw.append(["Показник","Значення"])
    dq=data.get("data_quality",{})
    for title,key in [("Групи дублів населених пунктів","settlement_duplicate_groups"),("Неканонічні населені пункти","settlement_noncanonical"),("Профілі без населеного пункту","profiles_without_settlement"),("Аномалії майбутньої відвідуваності","future_attendance_anomalies"),("Усього проблем","total_issues")]: dqw.append([title,dq.get(key,0)])
    dqw.column_dimensions["A"].width=44; dqw.column_dimensions["B"].width=18; _style_excel(dqw)

    dw=wb.create_sheet("Визначення")
    dw.append(["Показник","Визначення"])
    for title,description in data.get("indicator_definitions",[]): dw.append([title,description])
    dw.column_dimensions["A"].width=30; dw.column_dimensions["B"].width=95; _style_excel(dw)

    for title,key in [("Вік","age"),("Стать","gender"),("Населені пункти","settlement"),("Вразливість агреговано","vulnerability")]:
        sh=wb.create_sheet(title[:31]); sh.append(["Категорія","Кількість"])
        for lab,val in sorted(data[key].items(),key=lambda x:(-x[1],x[0])):
            shown = data.get("vulnerability_display", {}).get(lab, val) if key == "vulnerability" else val
            sh.append([lab,shown])
        sh.column_dimensions["A"].width=44; sh.column_dimensions["B"].width=15
        if sh.max_row>1:
            chart=BarChart(); chart.title=title; chart.add_data(Reference(sh,min_col=2,min_row=1,max_row=sh.max_row),titles_from_data=True); chart.set_categories(Reference(sh,min_col=1,min_row=2,max_row=sh.max_row)); chart.height=8; chart.width=15; sh.add_chart(chart,"D2")
        _style_excel(sh)
    iw=wb.create_sheet("Вплив")
    iw.append(["АМПасадори — результати / зворотний зв’язок", "Значення"])
    iw.append(["Показник", "Значення"])
    impact_rows=[
        ("Кількість завершених анкет зворотного зв’язку", data["outcomes"]["responses"]),
        ("Середня оцінка (1–5)", data["outcomes"]["avg_rating"]),
        ("Високо оцінили активності (4–5), %", data["outcomes"]["high_rating_pct"]),
        ("Було корисно, %", data["outcomes"]["useful_pct"]),
        ("Отримали нові знання, %", data["outcomes"]["new_knowledge_pct"]),
        ("Почувалися безпечно, %", data["outcomes"]["safe_pct"]),
        ("Хочуть прийти ще, %", data["outcomes"]["return_pct"]),
    ]
    for title,value in impact_rows: iw.append([title,value])
    iw.column_dimensions["A"].width=48; iw.column_dimensions["B"].width=20; _style_excel(iw)

    aw=wb.create_sheet("Розширена аналітика")
    aw.append(["АМПасадори — Розширена аналітика", "Значення"])
    aw.append(["Показник", "Значення"])
    for title,value in [
        ("Когорта — зареєструвалися",data["cohort_funnel"]["registered"]),
        ("Когорта — прийшли 1 раз",data["cohort_funnel"]["first_visit"]),
        ("Когорта — повернулися",data["cohort_funnel"]["returned"]),
        ("Когорта — стали регулярними",data["cohort_funnel"]["regular"]),
        ("Когорта — стали АМПасадорами",data["cohort_funnel"]["ambassadors"]),
        ("Повернення за 30 днів, %",data["retention"]["days30_pct"]),
        ("Повернення за 90 днів, %",data["retention"]["days90_pct"]),
        ("Середній індекс залученості",data["engagement"]["average"]),
        ("Залученість 75–100",data["engagement"]["high"]),
    ]: aw.append([title,value])
    aw.append([]); aw.append(["Теплова карта: день / година"]+[f"{h:02d}" for h in range(24)])
    for d,day in enumerate(data["heatmap_days"]): aw.append([day]+data["heatmap"][d])
    aw.column_dimensions["A"].width=38
    for c in range(2,26): aw.column_dimensions[get_column_letter(c)].width=6
    _style_excel(aw)

    gw=wb.create_sheet("Гейміфікація")
    gw.append(["АМПасадори — гейміфікація", "Значення"]); gw.append(["Показник","Значення"])
    gw.append(["Активна тижнева серія",data["streak_snapshot"]["weekly_active"]]); gw.append(["Активна суперсерія",data["streak_snapshot"]["super_active"]]); gw.append(["Серії, які можна відновити",data["streak_snapshot"]["recoverable"]]); gw.append(["Рекорд тижнів",data["streak_snapshot"]["best_weekly"]]); gw.append(["Рекорд подій",data["streak_snapshot"]["best_event"]]); gw.append(["Днів заморозки у вибраному періоді",data["summary"]["streak_freeze_days"]]); gw.append([]); gw.append(["Ліга","Учасники"])
    for title,count in data["league_distribution"].items(): gw.append([title,count])
    gw.column_dimensions["A"].width=42; gw.column_dimensions["B"].width=18; _style_excel(gw)

    bw=wb.create_sheet("Бейджі по тижнях"); bw.append(["Тиждень","Отримані бейджі"])
    for row in data["badge_weekly"]: bw.append([row["label"],row["value"]])
    bw.column_dimensions["A"].width=22; bw.column_dimensions["B"].width=20
    if bw.max_row>2:
        chart=LineChart(); chart.title="Отримані бейджі по тижнях"; chart.add_data(Reference(bw,min_col=2,min_row=1,max_row=bw.max_row),titles_from_data=True); chart.set_categories(Reference(bw,min_col=1,min_row=2,max_row=bw.max_row)); chart.height=8; chart.width=15; bw.add_chart(chart,"D2")
    elif bw.max_row==2:
        chart=BarChart(); chart.title="Отримані бейджі"; chart.add_data(Reference(bw,min_col=2,min_row=1,max_row=2),titles_from_data=True); chart.set_categories(Reference(bw,min_col=1,min_row=2,max_row=2)); chart.height=7; chart.width=12; bw.add_chart(chart,"D2")
    _style_excel(bw)

    bio=BytesIO(); wb.save(bio); return bio.getvalue()


def _plot_counter(ax, title: str, counter: Counter, *, items: list[tuple[str, Any]] | None = None) -> None:
    items=items if items is not None else sorted(counter.items(),key=lambda x:(-x[1],x[0]))
    if not items:
        ax.text(.5,.5,"Даних немає",ha="center",va="center"); ax.axis("off"); return
    labels=[x[0] for x in items]; vals=[x[1] for x in items]
    ax.barh(range(len(labels)),vals); ax.set_yticks(range(len(labels))); ax.set_yticklabels(labels); ax.invert_yaxis(); ax.set_title(title); ax.grid(axis="x",alpha=.2)


def report_pdf(data: dict[str, Any]) -> bytes:
    """Branded multi-page PDF. Never clips KPI cards or silently truncates tables."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.patches import FancyBboxPatch, Rectangle
    import textwrap

    plt.rcParams["font.family"]="DejaVu Sans"
    bio=BytesIO()
    logo_path=Path("app/web/static/amp_logo.png")
    logo=plt.imread(str(logo_path)) if logo_path.exists() else None
    brand="#0B5B6C"; accent="#06B8C5"; pale="#EAF8FA"; ink="#173B43"; muted="#6C858B"

    def wrap_words(text: str, width: int) -> str:
        return "\n".join(textwrap.wrap(str(text), width=width, break_long_words=False, break_on_hyphens=False))

    def footer(fig, right="Автоматизовано системою АМП XP"):
        fig.add_artist(Rectangle((.07,.055),.86,.002,transform=fig.transFigure,facecolor=accent,edgecolor="none"))
        fig.text(.07,.032,"АМПасадори • Анисівський молодіжний простір",fontsize=8,color=brand)
        fig.text(.93,.032,right,fontsize=7.5,color=muted,ha="right")

    def kpi_page(pdf, cards, *, title, subtitle="", first=False):
        fig=plt.figure(figsize=(8.27,11.69)); fig.patch.set_facecolor("white")
        if first:
            fig.add_artist(Rectangle((0,.82),1,.18,transform=fig.transFigure,facecolor=brand,edgecolor="none",zorder=-1))
            if logo is not None:
                axl=fig.add_axes([.065,.855,.14,.10]); axl.imshow(logo); axl.axis("off")
            fig.text(.225,.925,"АМПасадори",fontsize=24,weight="bold",color="white",va="center")
            fig.text(.225,.875,"Автоматичний звіт АМП",fontsize=16,weight="bold",color="#DDF8FA",va="center")
            fig.text(.07,.765,f"Період: {data['label']}",fontsize=12.5,weight="bold",color=ink)
            fig.text(.07,.735,f"Дані станом на {data['generated_at'].strftime('%d.%m.%Y %H:%M')} {data.get('timezone','Europe/Kyiv')}",fontsize=8.8,color=muted)
            fig.text(.07,.695,wrap_words(data["privacy_note"],105),fontsize=8.2,color=muted)
            top=.58
        else:
            fig.text(.07,.93,title,fontsize=22,weight="bold",color=brand)
            if subtitle:
                fig.text(.07,.875,wrap_words(subtitle,100),fontsize=10,color=muted)
                top=.72
            else:
                top=.80
        cols=3; card_w=.255; card_h=.105; gap_x=.025; gap_y=.024
        for idx,(label_text,value) in enumerate(cards):
            row=idx//cols; col=idx%cols; x=.07+col*(card_w+gap_x); y=top-row*(card_h+gap_y)
            card=FancyBboxPatch((x,y),card_w,card_h,boxstyle="round,pad=0.006,rounding_size=0.012",transform=fig.transFigure,facecolor=pale,edgecolor="#CFE5E9",linewidth=.8)
            fig.add_artist(card)
            fig.text(x+card_w/2,y+.069,wrap_words(label_text,24),fontsize=7.8,color=muted,ha="center",va="center")
            fig.text(x+card_w/2,y+.025,str(value),fontsize=15,weight="bold",color=brand,ha="center",va="center")
        footer(fig)
        pdf.savefig(fig); plt.close(fig)

    with PdfPages(bio) as pdf:
        period_labels=[("Завершені події","events"),("Усього заплановано","events_planned"),("Майбутні події","events_upcoming"),("Унікальні залучені","unique_participants"),("Підтверджені відвідування","visits"),("Середня відвідуваність","avg_attendance"),("Волонтерські години","volunteer_hours"),("Квести","quests_completed"),("Активності","activities_completed"),("Реалізовані ідеї","ideas_implemented"),("Звернення","requests"),("Нові учасники","new_participants"),("Нараховано XP","xp_awarded"),("Опитування","surveys_published"),("Відповіді на опитування","survey_responses"),("Отримані бейджі","badges_awarded"),("Дій участі","participation_actions"),("Середній XP/залученого","avg_xp_per_engaged"),("⚠ Аномалії відвідуваності","future_attendance_anomalies")]
        cards=[(title,data["period_summary"].get(key,0)) for title,key in period_labels]
        for idx in range(0,len(cards),12):
            kpi_page(pdf,cards[idx:idx+12],title="Потокові показники за період",subtitle="Потокові KPI: лише дії у вибраному періоді; у відвідуваності враховуються фактично завершені події" if idx==0 else "Продовження потокових показників",first=(idx==0))

        snapshot_labels=[("Активні профілі","active_profiles"),("Неактивні профілі","inactive_profiles"),("Видалені профілі","deleted_profiles"),("Видалені без відновлення","deleted_permanent_profiles"),("Запити на відновлення","restoration_requests_pending"),("На випробувальному строку","probation_profiles"),("Відновлені профілі","restored_profiles"),("Відхилені відновлення","restoration_rejected"),("Профілі з населеним пунктом","settlement_directory_profiles")]
        snapshot_cards=[(title,data["snapshot_summary"].get(key,0)) for title,key in snapshot_labels]
        kpi_page(pdf,snapshot_cards,title="Моментні показники станом на дату",subtitle=f"Дані станом на {data['generated_at'].strftime('%d.%m.%Y %H:%M')} {data.get('timezone','Europe/Kyiv')}")

        # Reporting 2.0 — conversion funnels.
        fig,axes=plt.subplots(1,2,figsize=(11.69,8.27)); fig.suptitle("Конверсійні воронки",fontsize=18,weight="bold")
        reg_order=[("Старт","start"),("Згода","consent"),("Профіль","profile"),("Надсилання","submit"),("Схвалення","approved"),("Перша активність","first_activity")]
        reg_vals=[data.get("registration_funnel",{}).get(key,0) for _,key in reg_order]
        axes[0].barh(range(len(reg_order)),reg_vals); axes[0].set_yticks(range(len(reg_order))); axes[0].set_yticklabels([x[0] for x in reg_order]); axes[0].invert_yaxis(); axes[0].set_title("Реєстрація"); axes[0].grid(axis="x",alpha=.2)
        event_order=[("Заявки","registered"),("Відмітка","checkin"),("Підтверджено","attended"),("XP","xp"),("Відгук","feedback")]
        event_vals=[data.get("event_conversion",{}).get(key,0) for _,key in event_order]
        axes[1].barh(range(len(event_order)),event_vals); axes[1].set_yticks(range(len(event_order))); axes[1].set_yticklabels([x[0] for x in event_order]); axes[1].invert_yaxis(); axes[1].set_title("Участь у подіях"); axes[1].grid(axis="x",alpha=.2)
        fig.text(.06,.04,"Воронка реєстрації — когорта, що стартувала у вибраному періоді. Воронка подій — реєстрації на події вибраного періоду.",fontsize=8.5,color=muted)
        fig.tight_layout(rect=[0,.07,1,.92]); pdf.savefig(fig); plt.close(fig)

        dq=data.get("data_quality",{})
        quality_cards=[
            ("Групи дублів населених пунктів",dq.get("settlement_duplicate_groups",0)),
            ("Неканонічні населені пункти",dq.get("settlement_noncanonical",0)),
            ("Профілі без населеного пункту",dq.get("profiles_without_settlement",0)),
            ("Аномалії майбутньої відвідуваності",dq.get("future_attendance_anomalies",0)),
            ("Усього проблем",dq.get("total_issues",0)),
        ]
        kpi_page(pdf,quality_cards,title="Якість даних",subtitle="Блок якості даних: проблеми, які можуть спотворювати сегментацію або звітність")

        fig=plt.figure(figsize=(8.27,11.69)); fig.patch.set_facecolor("white")
        fig.text(.07,.93,"Визначення показників",fontsize=22,weight="bold",color=brand)
        fig.text(.07,.89,"Що саме означають основні KPI цього звіту",fontsize=10,color=muted)
        y=.83
        for title_text,description in data.get("indicator_definitions",[]):
            fig.text(.08,y,title_text,fontsize=10.5,weight="bold",color=ink,va="top")
            y-=.028
            fig.text(.08,y,wrap_words(description,100),fontsize=8.7,color=muted,va="top",linespacing=1.35)
            y-=.075
        footer(fig); pdf.savefig(fig); plt.close(fig)

        advanced=[
            ("Прийшли 1 раз",data["cohort_funnel"]["first_visit"]),
            ("Повернулися",data["cohort_funnel"]["returned"]),
            ("Стали регулярними",data["cohort_funnel"]["regular"]),
            ("Стали АМПасадорами",data["cohort_funnel"]["ambassadors"]),
            ("Повернення за 30 днів",f"{data['retention']['days30_pct']:g}%"),
            ("Повернення за 90 днів",f"{data['retention']['days90_pct']:g}%"),
            ("Індекс залученості",f"{data['engagement']['average']:g}/100"),
            ("Залученість 75–100",data["engagement"]["high"]),
        ]
        kpi_page(pdf,advanced,title="Розширена аналітика",subtitle="Когортний аналіз, повернення та внутрішній індекс залученості")

        # Outcomes page.
        fig=plt.figure(figsize=(8.27,11.69)); fig.patch.set_facecolor("white")
        fig.text(.07,.93,"Вплив",fontsize=24,weight="bold",color=brand)
        fig.text(.07,.885,"Зворотний зв’язок учасників після подій",fontsize=12,color=muted)
        o=data.get("outcomes",{})
        impact=[
            ("Високо оцінили активності",f"{o.get('high_rating_pct',0):g}%"),
            ("Отримали нові знання",f"{o.get('new_knowledge_pct',0):g}%"),
            ("Почувалися безпечно",f"{o.get('safe_pct',0):g}%"),
            ("Вважають активності корисними",f"{o.get('useful_pct',0):g}%"),
            ("Хочуть прийти ще",f"{o.get('return_pct',0):g}%"),
            ("Середня оцінка",f"{o.get('avg_rating',0):g}/5"),
        ]
        if int(o.get("responses",0))>0:
            fig.text(.07,.835,f"{o.get('high_rating_pct',0):g}% учасників високо оцінили активності.",fontsize=9.5,color=muted)
            fig.text(.07,.812,f"{o.get('new_knowledge_pct',0):g}% повідомили про отримання нових знань; {o.get('safe_pct',0):g}% почувалися безпечно.",fontsize=9.5,color=muted)
            y=.74
            for title,value in impact:
                fig.text(.08,y,wrap_words(title,48),fontsize=11,color=ink,va="center")
                fig.text(.88,y,value,fontsize=15,weight="bold",color=brand,ha="right",va="center")
                y-=.085
            fig.text(.08,.25,f"Завершених анкет зворотного зв’язку: {o.get('responses',0)}",fontsize=10,color=muted)
        else:
            fig.text(.07,.78,"У вибраному періоді завершених feedback-анкет ще немає.",fontsize=14,color=muted)
        fig.text(.07,.15,wrap_words("Показники впливу доповнюють кількісні дані та показують сприйняту користь, навчальний результат, безпеку й намір повернутися.",105),fontsize=9,color=muted)
        footer(fig); pdf.savefig(fig); plt.close(fig)

        # Activity heatmap.
        fig,ax=plt.subplots(figsize=(11.69,8.27)); matrix=data.get("heatmap") or [[0]*24 for _ in range(7)]
        im=ax.imshow(matrix,aspect="auto",interpolation="nearest",cmap="YlGnBu")
        ax.set_yticks(range(7)); ax.set_yticklabels(data.get("heatmap_days",["Пн","Вт","Ср","Чт","Пт","Сб","Нд"]))
        ax.set_xticks(range(24)); ax.set_xticklabels([f"{h:02d}" for h in range(24)],fontsize=8)
        ax.set_xlabel("Година"); ax.set_ylabel("День тижня"); ax.set_title("Теплова карта активності: день тижня × година",fontsize=16,weight="bold",pad=18)
        fig.colorbar(im,ax=ax,fraction=.025,pad=.02,label="Кількість дій")
        fig.tight_layout(); pdf.savefig(fig); plt.close(fig)

        # Reporting 2.0 dynamics: daily for short periods, weekly for medium
        # periods, monthly for long periods. Never draw a meaningless one-point line.
        rows=data.get("trend",data.get("monthly",[]))
        if len(rows)>1:
            fig,ax=plt.subplots(figsize=(11.69,8.27)); x=range(len(rows))
            ax.plot(x,[r["visits"] for r in rows],marker="o",label="Відвідування",color=brand,linewidth=2.2)
            ax.plot(x,[r["new_users"] for r in rows],marker="o",label="Нові учасники",color=accent,linewidth=2.2)
            ax.set_xticks(list(x)); ax.set_xticklabels([r["label"] for r in rows],rotation=35,ha="right"); ax.legend(); ax.grid(axis="y",alpha=.2)
            ax.set_title(f"Динаміка за період {data.get('trend_granularity_label','')}",fontsize=16,weight="bold"); fig.tight_layout(); pdf.savefig(fig); plt.close(fig)
        elif len(rows)==1:
            r=rows[0]
            kpi_page(pdf,[("Відвідування",r["visits"]),("Нові учасники",r["new_users"]),("Завершені події",r["events"]),("Нараховано XP",r["xp"]),("Отримані бейджі",r.get("badges",0))],title="Динаміка за період",subtitle="Є лише одна точка даних — замість лінійного графіка показано KPI.")
        else:
            kpi_page(pdf,[("Відвідування",0),("Нові учасники",0),("Завершені події",0),("Нараховано XP",0)],title="Динаміка за період",subtitle="У вибраному періоді немає даних для побудови динаміки.")

        # Events: all rows, paginated without truncation.
        event_rows=data["events"]
        if not event_rows:
            fig,ax=plt.subplots(figsize=(11.69,8.27)); ax.axis("off"); ax.set_title("Події та відвідуваність",fontsize=16,weight="bold",pad=20); ax.text(.5,.5,"Подій у вибраному періоді немає",ha="center",va="center"); pdf.savefig(fig); plt.close(fig)
        else:
            for page_no,start in enumerate(range(0,len(event_rows),20),1):
                chunk=event_rows[start:start+20]
                fig,ax=plt.subplots(figsize=(11.69,8.27)); ax.axis("off"); ax.set_title(f"Події та відвідуваність — сторінка {page_no}",fontsize=16,weight="bold",pad=20)
                rows=[[r["date"],wrap_words(r["title"],38),r.get("timing_label",r["status"]),str(r["registered"]),str(r["attended"])] for r in chunk]
                table=ax.table(cellText=rows,colLabels=["Дата","Подія","Стан","Зареєстровано","Відвідали"],loc="upper center",cellLoc="left",colWidths=[.17,.43,.14,.13,.13]); table.auto_set_font_size(False); table.set_fontsize(8.2); table.scale(1,1.55)
                for (rr,cc),cell in table.get_celld().items():
                    cell.set_edgecolor("#D4E6E9"); cell.set_linewidth(.55); cell.get_text().set_wrap(True)
                    if rr==0: cell.set_facecolor(brand); cell.set_text_props(color="white",weight="bold",ha="center")
                    elif rr%2==0: cell.set_facecolor("#F7FBFC")
                fig.tight_layout(); pdf.savefig(fig); plt.close(fig)

        # Demographics: paginate every 15 rows so no category disappears.
        for title,key in [("Вікова статистика","age"),("Гендерна статистика","gender"),("Населені пункти","settlement"),("Категорії вразливості — тільки агреговано","vulnerability")]:
            all_items=sorted(data[key].items(),key=lambda x:(-x[1],x[0]))
            chunks=[all_items[i:i+15] for i in range(0,len(all_items),15)] or [[]]
            for page_no,chunk in enumerate(chunks,1):
                fig,ax=plt.subplots(figsize=(11.69,8.27)); _plot_counter(ax,title if len(chunks)==1 else f"{title} — {page_no}/{len(chunks)}",data[key],items=chunk)
                if key=="vulnerability": fig.text(.08,.04,wrap_words(data["privacy_note"],140),fontsize=8)
                fig.tight_layout(rect=[0,.06 if key=="vulnerability" else 0,1,1]); pdf.savefig(fig); plt.close(fig)

        fig,axes=plt.subplots(1,2,figsize=(11.69,8.27)); fig.suptitle("Гейміфікація: ліги та серії",fontsize=16,weight="bold")
        leagues=list(data["league_distribution"].items()); axes[0].barh(range(len(leagues)),[v for _,v in leagues]); axes[0].set_yticks(range(len(leagues))); axes[0].set_yticklabels([k for k,_ in leagues]); axes[0].invert_yaxis(); axes[0].set_title("Ліги — поточний стан"); axes[0].grid(axis="x",alpha=.2)
        ss=data["streak_snapshot"]; streak_items=[("Тижнева",ss["weekly_active"]),("Суперсерія",ss["super_active"]),("Можна відновити",ss["recoverable"])]; axes[1].bar([x[0] for x in streak_items],[x[1] for x in streak_items]); axes[1].set_title("Серії участі"); axes[1].tick_params(axis="x",rotation=20); axes[1].grid(axis="y",alpha=.2); fig.text(.06,.04,f"Заморозка серій у періоді: {data['summary']['streak_freeze_days']} дн.",fontsize=9); fig.tight_layout(rect=[0,.06,1,.93]); pdf.savefig(fig); plt.close(fig)

        fig,ax=plt.subplots(figsize=(11.69,8.27)); rows=data["badge_weekly"]
        if len(rows)>1:
            x=range(len(rows)); ax.plot(x,[r["value"] for r in rows],marker="o"); ax.set_xticks(list(x)); ax.set_xticklabels([r["label"] for r in rows],rotation=35,ha="right"); ax.grid(axis="y",alpha=.2)
        elif len(rows)==1:
            ax.bar([rows[0]["label"]],[rows[0]["value"]]); ax.grid(axis="y",alpha=.2)
        else: ax.text(.5,.5,"У вибраному періоді бейджів не видавали",ha="center",va="center")
        ax.set_title("Отримані бейджі по тижнях",fontsize=16,weight="bold"); fig.tight_layout(); pdf.savefig(fig); plt.close(fig)
    return bio.getvalue()

