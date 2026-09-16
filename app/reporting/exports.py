from __future__ import annotations

from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Any

def _style_excel(ws) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    teal="06AEBB"; deep="0B5B6C"; pale="EAF8FA"; alt="F7FBFC"; dark="173B43"; muted="6C858B"; white="FFFFFF"
    thin=Side(style="thin",color="D7E5E8")
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.tabColor=teal
    ws.page_margins.left=.35; ws.page_margins.right=.35; ws.page_margins.top=.55; ws.page_margins.bottom=.55
    ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.oddHeader.center.text="&BАМПасадори • Анисівський молодіжний простір"
    ws.oddHeader.center.size=10
    ws.oddHeader.center.font="Arial,Bold"
    ws.oddFooter.left.text="АМП • автоматичний звіт"
    ws.oddFooter.right.text="Сторінка &P з &N"
    for row in ws.iter_rows():
        for c in row:
            c.alignment=Alignment(vertical="top",wrap_text=True)
            c.border=Border(bottom=thin)
            c.font=Font(name="Arial",size=10,color=dark)
    # Brand the first row as a title/header row.
    for c in ws[1]:
        c.font=Font(name="Arial",bold=True,size=14,color=white)
        c.fill=PatternFill("solid",fgColor=deep)
        c.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[1].height=26
    # Detect secondary table headers used by the report sheets.
    header_names={"Показник","Подія","Місяць","Категорія","Ліга","Тиждень"}
    for row in ws.iter_rows(min_row=2):
        first=str(row[0].value or "")
        if first in header_names:
            for c in row:
                c.font=Font(name="Arial",bold=True,size=10,color=deep)
                c.fill=PatternFill("solid",fgColor=pale)
                c.alignment=Alignment(vertical="center",wrap_text=True)
        elif row[0].row % 2 == 0:
            for c in row:
                if c.value is not None and c.fill.fill_type is None:
                    c.fill=PatternFill("solid",fgColor=alt)
    ws.freeze_panes="A2"

def report_excel(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter
    wb=Workbook(); ws=wb.active; ws.title="Зведення"
    ws.append(["АМПасадори — автоматичний звіт",""])
    ws.append(["Період",data["label"]]); ws.append(["Дані станом на",f"{data['generated_at'].strftime('%d.%m.%Y %H:%M')} {data.get('timezone','Europe/Kyiv')}"]); ws.append(["Примітка",data["privacy_note"]]); ws.append([])
    labels={"events":"Завершені події","events_planned":"Усього заплановано подій","events_upcoming":"Майбутні події","events_in_progress":"Події у поточному вікні","unique_participants":"Унікальні залучені учасники","visits":"Підтверджені відвідування","avg_attendance":"Середня відвідуваність","volunteer_hours":"Волонтерські години","tasks_completed":"Виконані волонтерські задачі","quests_completed":"Підтверджені квести","activities_completed":"Підтверджені активності","ideas_submitted":"Подані ідеї","ideas_implemented":"Реалізовані ідеї","requests":"Звернення","requests_resolved":"Вирішені звернення","new_participants":"Нові учасники","xp_awarded":"Нараховано XP","opportunity_interests":"Позначки «Мені цікаво»","surveys_published":"Опубліковані опитування","survey_responses":"Проходження опитувань","badges_awarded":"Отримані бейджі","streak_freeze_days":"Днів заморозки серій","active_profiles":"Активні профілі","inactive_profiles":"Неактивні профілі","deleted_profiles":"Видалені профілі","deleted_permanent_profiles":"Видалені без відновлення","restoration_requests_pending":"Запити на відновлення","probation_profiles":"На 14-денному випробувальному строку","restored_profiles":"Відновлені профілі","restoration_rejected":"Відхилені запити на відновлення","participation_actions":"Дій участі","avg_xp_per_engaged":"Середній XP на залученого","future_attendance_anomalies":"Некоректні підтвердження участі поза завершеними подіями","settlement_directory_profiles":"Профілі з канонічним населеним пунктом","feedback_responses":"Зворотний зв’язок — відповідей","feedback_avg_rating":"Зворотний зв’язок — середня оцінка","feedback_high_rating_pct":"Висока оцінка 4–5, %","feedback_useful_pct":"Було корисно, %","feedback_new_knowledge_pct":"Нові знання, %","feedback_safe_pct":"Почувалися безпечно, %","feedback_return_pct":"Хочуть прийти ще, %","cohort_first_visit":"Когорта — прийшли 1 раз","cohort_returned":"Когорта — повернулися","cohort_regular":"Когорта — регулярні","cohort_ambassadors":"Когорта — АМПасадори","retention_30_pct":"Повернення за 30 днів, %","retention_90_pct":"Повернення за 90 днів, %","engagement_average":"Середній індекс залученості","engagement_high":"Залученість 75–100"}
    ws.append(["ПОТОКОВІ ПОКАЗНИКИ ЗА ПЕРІОД", ""]); ws.append(["Показник","Значення"])
    for k,v in data.get("period_summary", data["summary"]).items(): ws.append([labels.get(k,k),v])
    ws.append([]); ws.append(["МОМЕНТНІ ПОКАЗНИКИ СТАНОМ НА ДАТУ", f"{data['generated_at'].strftime('%d.%m.%Y %H:%M')} {data.get('timezone','Europe/Kyiv')}"]); ws.append(["Показник","Значення"])
    for k,v in data.get("snapshot_summary", {}).items(): ws.append([labels.get(k,k),v])
    ws.column_dimensions["A"].width=44; ws.column_dimensions["B"].width=24; _style_excel(ws)

    ew=wb.create_sheet("Події"); ew.append(["Подія","Дата/час","Локація","Стан у звіті","Системний статус","Зареєстровано","Відвідали"])
    for r in data["events"]: ew.append([r["title"],r["date"],r["location"],r.get("timing_label",r["status"]),r["status"],r["registered"],r["attended"]])
    for i,w in enumerate([42,20,26,18,16,16,14],1): ew.column_dimensions[get_column_letter(i)].width=w
    _style_excel(ew)

    mw=wb.create_sheet("Динаміка"); mw.append([f"Період ({data.get('trend_granularity_label','')})","Завершені події","Заплановані події","Відвідування","Нові учасники","XP","Опитування — відповіді","Отримані бейджі"])
    trend_rows=data.get("trend",data.get("monthly",[]))
    for r in trend_rows: mw.append([r["label"],r["events"],r.get("planned_events",r["events"]),r["visits"],r["new_users"],r["xp"],r.get("survey_responses",0),r.get("badges",0)])
    if len(trend_rows)>1:
        chart=LineChart(); chart.title=f"Динаміка відвідувань {data.get('trend_granularity_label','')}"; chart.add_data(Reference(mw,min_col=4,min_row=1,max_row=mw.max_row),titles_from_data=True); chart.set_categories(Reference(mw,min_col=1,min_row=2,max_row=mw.max_row)); mw.add_chart(chart,"J2")
    elif len(trend_rows)==1:
        mw["J2"]="Одна точка даних — лінійний графік не будується."
        mw["J3"]="Відвідування"; mw["K3"]=trend_rows[0]["visits"]
        mw["J4"]="Нові учасники"; mw["K4"]=trend_rows[0]["new_users"]
        mw["J5"]="XP"; mw["K5"]=trend_rows[0]["xp"]
    _style_excel(mw)

    fw=wb.create_sheet("Воронки")
    fw.append(["Конверсія реєстрації", "Кількість"]); fw.append(["Етап","Кількість"])
    reg_labels={"start":"Старт","consent":"Згода","profile":"Профіль","submit":"Надсилання","approved":"Схвалення","first_activity":"Перша активність"}
    for key in ["start","consent","profile","submit","approved","first_activity"]: fw.append([reg_labels[key],data.get("registration_funnel",{}).get(key,0)])
    fw.append([]); fw.append(["Конверсія участі у подіях", "Кількість"]); fw.append(["Етап","Кількість"])
    event_labels={"registered":"Заявки / реєстрації","checkin":"Відмітка","attended":"Підтверджена участь","xp":"XP нараховано","feedback":"Завершений відгук"}
    for key in ["registered","checkin","attended","xp","feedback"]: fw.append([event_labels[key],data.get("event_conversion",{}).get(key,0)])
    fw.column_dimensions["A"].width=34; fw.column_dimensions["B"].width=18; _style_excel(fw)

    dqw=wb.create_sheet("Якість даних")
    dqw.append(["Якість даних", "Кількість"]); dqw.append(["Показник","Значення"])
    dq=data.get("data_quality",{})
    for title,key in [("Групи дублів населених пунктів","settlement_duplicate_groups"),("Неканонічні населені пункти","settlement_noncanonical"),("Профілі без населеного пункту","profiles_without_settlement"),("Аномалії майбутньої відвідуваності","future_attendance_anomalies"),("Усього проблем","total_issues")]: dqw.append([title,dq.get(key,0)])
    dqw.column_dimensions["A"].width=44; dqw.column_dimensions["B"].width=18; _style_excel(dqw)

    dw=wb.create_sheet("Визначення")
    dw.append(["Показник","Визначення"])
    for title,description in data.get("indicator_definitions",[]): dw.append([title,description])
    dw.column_dimensions["A"].width=30; dw.column_dimensions["B"].width=95; _style_excel(dw)

    for title,key in [("Вік","age"),("Стать","gender"),("Населені пункти","settlement"),("Вразливість агреговано","vulnerability")]:
        sh=wb.create_sheet(title[:31]); sh.append(["Категорія","Кількість"])
        for lab,val in sorted(data[key].items(),key=lambda x:(-x[1],x[0])):
            shown = data.get("vulnerability_display", {}).get(lab, val) if key == "vulnerability" else val
            sh.append([lab,shown])
        sh.column_dimensions["A"].width=44; sh.column_dimensions["B"].width=15
        if sh.max_row>1:
            chart=BarChart(); chart.title=title; chart.add_data(Reference(sh,min_col=2,min_row=1,max_row=sh.max_row),titles_from_data=True); chart.set_categories(Reference(sh,min_col=1,min_row=2,max_row=sh.max_row)); chart.height=8; chart.width=15; sh.add_chart(chart,"D2")
        _style_excel(sh)
    iw=wb.create_sheet("Вплив")
    iw.append(["АМПасадори — результати / зворотний зв’язок", "Значення"])
    iw.append(["Показник", "Значення"])
    impact_rows=[
        ("Кількість завершених анкет зворотного зв’язку", data["outcomes"]["responses"]),
        ("Середня оцінка (1–5)", data["outcomes"]["avg_rating"]),
        ("Високо оцінили активності (4–5), %", data["outcomes"]["high_rating_pct"]),
        ("Було корисно, %", data["outcomes"]["useful_pct"]),
        ("Отримали нові знання, %", data["outcomes"]["new_knowledge_pct"]),
        ("Почувалися безпечно, %", data["outcomes"]["safe_pct"]),
        ("Хочуть прийти ще, %", data["outcomes"]["return_pct"]),
    ]
    for title,value in impact_rows: iw.append([title,value])
    iw.column_dimensions["A"].width=48; iw.column_dimensions["B"].width=20; _style_excel(iw)

    aw=wb.create_sheet("Розширена аналітика")
    aw.append(["АМПасадори — Розширена аналітика", "Значення"])
    aw.append(["Показник", "Значення"])
    for title,value in [
        ("Когорта — зареєструвалися",data["cohort_funnel"]["registered"]),
        ("Когорта — прийшли 1 раз",data["cohort_funnel"]["first_visit"]),
        ("Когорта — повернулися",data["cohort_funnel"]["returned"]),
        ("Когорта — стали регулярними",data["cohort_funnel"]["regular"]),
        ("Когорта — стали АМПасадорами",data["cohort_funnel"]["ambassadors"]),
        ("Повернення за 30 днів, %",data["retention"]["days30_pct"]),
        ("Повернення за 90 днів, %",data["retention"]["days90_pct"]),
        ("Середній індекс залученості",data["engagement"]["average"]),
        ("Залученість 75–100",data["engagement"]["high"]),
    ]: aw.append([title,value])
    aw.append([]); aw.append(["Теплова карта: день / година"]+[f"{h:02d}" for h in range(24)])
    for d,day in enumerate(data["heatmap_days"]): aw.append([day]+data["heatmap"][d])
    aw.column_dimensions["A"].width=38
    for c in range(2,26): aw.column_dimensions[get_column_letter(c)].width=6
    _style_excel(aw)

    gw=wb.create_sheet("Гейміфікація")
    gw.append(["АМПасадори — гейміфікація", "Значення"]); gw.append(["Показник","Значення"])
    gw.append(["Активна тижнева серія",data["streak_snapshot"]["weekly_active"]]); gw.append(["Активна суперсерія",data["streak_snapshot"]["super_active"]]); gw.append(["Серії, які можна відновити",data["streak_snapshot"]["recoverable"]]); gw.append(["Рекорд тижнів",data["streak_snapshot"]["best_weekly"]]); gw.append(["Рекорд подій",data["streak_snapshot"]["best_event"]]); gw.append(["Днів заморозки у вибраному періоді",data["summary"]["streak_freeze_days"]]); gw.append([]); gw.append(["Ліга","Учасники"])
    for title,count in data["league_distribution"].items(): gw.append([title,count])
    gw.column_dimensions["A"].width=42; gw.column_dimensions["B"].width=18; _style_excel(gw)

    bw=wb.create_sheet("Бейджі по тижнях"); bw.append(["Тиждень","Отримані бейджі"])
    for row in data["badge_weekly"]: bw.append([row["label"],row["value"]])
    bw.column_dimensions["A"].width=22; bw.column_dimensions["B"].width=20
    if bw.max_row>2:
        chart=LineChart(); chart.title="Отримані бейджі по тижнях"; chart.add_data(Reference(bw,min_col=2,min_row=1,max_row=bw.max_row),titles_from_data=True); chart.set_categories(Reference(bw,min_col=1,min_row=2,max_row=bw.max_row)); chart.height=8; chart.width=15; bw.add_chart(chart,"D2")
    elif bw.max_row==2:
        chart=BarChart(); chart.title="Отримані бейджі"; chart.add_data(Reference(bw,min_col=2,min_row=1,max_row=2),titles_from_data=True); chart.set_categories(Reference(bw,min_col=1,min_row=2,max_row=2)); chart.height=7; chart.width=12; bw.add_chart(chart,"D2")
    _style_excel(bw)

    bio=BytesIO(); wb.save(bio); return bio.getvalue()

def _plot_counter(ax, title: str, counter: Counter, *, items: list[tuple[str, Any]] | None = None) -> None:
    items=items if items is not None else sorted(counter.items(),key=lambda x:(-x[1],x[0]))
    if not items:
        ax.text(.5,.5,"Даних немає",ha="center",va="center"); ax.axis("off"); return
    labels=[x[0] for x in items]; vals=[x[1] for x in items]
    ax.barh(range(len(labels)),vals); ax.set_yticks(range(len(labels))); ax.set_yticklabels(labels); ax.invert_yaxis(); ax.set_title(title); ax.grid(axis="x",alpha=.2)

def report_pdf(data: dict[str, Any]) -> bytes:
    """Branded multi-page PDF. Never clips KPI cards or silently truncates tables."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.patches import FancyBboxPatch, Rectangle
    import textwrap

    plt.rcParams["font.family"]="DejaVu Sans"
    bio=BytesIO()
    logo_path=Path("app/web/static/amp_logo.png")
    logo=plt.imread(str(logo_path)) if logo_path.exists() else None
    brand="#0B5B6C"; accent="#06B8C5"; pale="#EAF8FA"; ink="#173B43"; muted="#6C858B"

    def wrap_words(text: str, width: int) -> str:
        return "\n".join(textwrap.wrap(str(text), width=width, break_long_words=False, break_on_hyphens=False))

    def footer(fig, right="Автоматизовано системою АМП XP"):
        fig.add_artist(Rectangle((.07,.055),.86,.002,transform=fig.transFigure,facecolor=accent,edgecolor="none"))
        fig.text(.07,.032,"АМПасадори • Анисівський молодіжний простір",fontsize=8,color=brand)
        fig.text(.93,.032,right,fontsize=7.5,color=muted,ha="right")

    def kpi_page(pdf, cards, *, title, subtitle="", first=False):
        fig=plt.figure(figsize=(8.27,11.69)); fig.patch.set_facecolor("white")
        if first:
            fig.add_artist(Rectangle((0,.82),1,.18,transform=fig.transFigure,facecolor=brand,edgecolor="none",zorder=-1))
            if logo is not None:
                axl=fig.add_axes([.065,.855,.14,.10]); axl.imshow(logo); axl.axis("off")
            fig.text(.225,.925,"АМПасадори",fontsize=24,weight="bold",color="white",va="center")
            fig.text(.225,.875,"Автоматичний звіт АМП",fontsize=16,weight="bold",color="#DDF8FA",va="center")
            fig.text(.07,.765,f"Період: {data['label']}",fontsize=12.5,weight="bold",color=ink)
            fig.text(.07,.735,f"Дані станом на {data['generated_at'].strftime('%d.%m.%Y %H:%M')} {data.get('timezone','Europe/Kyiv')}",fontsize=8.8,color=muted)
            fig.text(.07,.695,wrap_words(data["privacy_note"],105),fontsize=8.2,color=muted)
            top=.58
        else:
            fig.text(.07,.93,title,fontsize=22,weight="bold",color=brand)
            if subtitle:
                fig.text(.07,.875,wrap_words(subtitle,100),fontsize=10,color=muted)
                top=.72
            else:
                top=.80
        cols=3; card_w=.255; card_h=.105; gap_x=.025; gap_y=.024
        for idx,(label_text,value) in enumerate(cards):
            row=idx//cols; col=idx%cols; x=.07+col*(card_w+gap_x); y=top-row*(card_h+gap_y)
            card=FancyBboxPatch((x,y),card_w,card_h,boxstyle="round,pad=0.006,rounding_size=0.012",transform=fig.transFigure,facecolor=pale,edgecolor="#CFE5E9",linewidth=.8)
            fig.add_artist(card)
            fig.text(x+card_w/2,y+.069,wrap_words(label_text,24),fontsize=7.8,color=muted,ha="center",va="center")
            fig.text(x+card_w/2,y+.025,str(value),fontsize=15,weight="bold",color=brand,ha="center",va="center")
        footer(fig)
        pdf.savefig(fig); plt.close(fig)

    with PdfPages(bio) as pdf:
        period_labels=[("Завершені події","events"),("Усього заплановано","events_planned"),("Майбутні події","events_upcoming"),("Унікальні залучені","unique_participants"),("Підтверджені відвідування","visits"),("Середня відвідуваність","avg_attendance"),("Волонтерські години","volunteer_hours"),("Квести","quests_completed"),("Активності","activities_completed"),("Реалізовані ідеї","ideas_implemented"),("Звернення","requests"),("Нові учасники","new_participants"),("Нараховано XP","xp_awarded"),("Опитування","surveys_published"),("Відповіді на опитування","survey_responses"),("Отримані бейджі","badges_awarded"),("Дій участі","participation_actions"),("Середній XP/залученого","avg_xp_per_engaged"),("⚠ Аномалії відвідуваності","future_attendance_anomalies")]
        cards=[(title,data["period_summary"].get(key,0)) for title,key in period_labels]
        for idx in range(0,len(cards),12):
            kpi_page(pdf,cards[idx:idx+12],title="Потокові показники за період",subtitle="Потокові KPI: лише дії у вибраному періоді; у відвідуваності враховуються фактично завершені події" if idx==0 else "Продовження потокових показників",first=(idx==0))

        snapshot_labels=[("Активні профілі","active_profiles"),("Неактивні профілі","inactive_profiles"),("Видалені профілі","deleted_profiles"),("Видалені без відновлення","deleted_permanent_profiles"),("Запити на відновлення","restoration_requests_pending"),("На випробувальному строку","probation_profiles"),("Відновлені профілі","restored_profiles"),("Відхилені відновлення","restoration_rejected"),("Профілі з населеним пунктом","settlement_directory_profiles")]
        snapshot_cards=[(title,data["snapshot_summary"].get(key,0)) for title,key in snapshot_labels]
        kpi_page(pdf,snapshot_cards,title="Моментні показники станом на дату",subtitle=f"Дані станом на {data['generated_at'].strftime('%d.%m.%Y %H:%M')} {data.get('timezone','Europe/Kyiv')}")

        # Reporting 2.0 — conversion funnels.
        fig,axes=plt.subplots(1,2,figsize=(11.69,8.27)); fig.suptitle("Конверсійні воронки",fontsize=18,weight="bold")
        reg_order=[("Старт","start"),("Згода","consent"),("Профіль","profile"),("Надсилання","submit"),("Схвалення","approved"),("Перша активність","first_activity")]
        reg_vals=[data.get("registration_funnel",{}).get(key,0) for _,key in reg_order]
        axes[0].barh(range(len(reg_order)),reg_vals); axes[0].set_yticks(range(len(reg_order))); axes[0].set_yticklabels([x[0] for x in reg_order]); axes[0].invert_yaxis(); axes[0].set_title("Реєстрація"); axes[0].grid(axis="x",alpha=.2)
        event_order=[("Заявки","registered"),("Відмітка","checkin"),("Підтверджено","attended"),("XP","xp"),("Відгук","feedback")]
        event_vals=[data.get("event_conversion",{}).get(key,0) for _,key in event_order]
        axes[1].barh(range(len(event_order)),event_vals); axes[1].set_yticks(range(len(event_order))); axes[1].set_yticklabels([x[0] for x in event_order]); axes[1].invert_yaxis(); axes[1].set_title("Участь у подіях"); axes[1].grid(axis="x",alpha=.2)
        fig.text(.06,.04,"Воронка реєстрації — когорта, що стартувала у вибраному періоді. Воронка подій — реєстрації на події вибраного періоду.",fontsize=8.5,color=muted)
        fig.tight_layout(rect=[0,.07,1,.92]); pdf.savefig(fig); plt.close(fig)

        dq=data.get("data_quality",{})
        quality_cards=[
            ("Групи дублів населених пунктів",dq.get("settlement_duplicate_groups",0)),
            ("Неканонічні населені пункти",dq.get("settlement_noncanonical",0)),
            ("Профілі без населеного пункту",dq.get("profiles_without_settlement",0)),
            ("Аномалії майбутньої відвідуваності",dq.get("future_attendance_anomalies",0)),
            ("Усього проблем",dq.get("total_issues",0)),
        ]
        kpi_page(pdf,quality_cards,title="Якість даних",subtitle="Блок якості даних: проблеми, які можуть спотворювати сегментацію або звітність")

        fig=plt.figure(figsize=(8.27,11.69)); fig.patch.set_facecolor("white")
        fig.text(.07,.93,"Визначення показників",fontsize=22,weight="bold",color=brand)
        fig.text(.07,.89,"Що саме означають основні KPI цього звіту",fontsize=10,color=muted)
        y=.83
        for title_text,description in data.get("indicator_definitions",[]):
            fig.text(.08,y,title_text,fontsize=10.5,weight="bold",color=ink,va="top")
            y-=.028
            fig.text(.08,y,wrap_words(description,100),fontsize=8.7,color=muted,va="top",linespacing=1.35)
            y-=.075
        footer(fig); pdf.savefig(fig); plt.close(fig)

        advanced=[
            ("Прийшли 1 раз",data["cohort_funnel"]["first_visit"]),
            ("Повернулися",data["cohort_funnel"]["returned"]),
            ("Стали регулярними",data["cohort_funnel"]["regular"]),
            ("Стали АМПасадорами",data["cohort_funnel"]["ambassadors"]),
            ("Повернення за 30 днів",f"{data['retention']['days30_pct']:g}%"),
            ("Повернення за 90 днів",f"{data['retention']['days90_pct']:g}%"),
            ("Індекс залученості",f"{data['engagement']['average']:g}/100"),
            ("Залученість 75–100",data["engagement"]["high"]),
        ]
        kpi_page(pdf,advanced,title="Розширена аналітика",subtitle="Когортний аналіз, повернення та внутрішній індекс залученості")

        # Outcomes page.
        fig=plt.figure(figsize=(8.27,11.69)); fig.patch.set_facecolor("white")
        fig.text(.07,.93,"Вплив",fontsize=24,weight="bold",color=brand)
        fig.text(.07,.885,"Зворотний зв’язок учасників після подій",fontsize=12,color=muted)
        o=data.get("outcomes",{})
        impact=[
            ("Високо оцінили активності",f"{o.get('high_rating_pct',0):g}%"),
            ("Отримали нові знання",f"{o.get('new_knowledge_pct',0):g}%"),
            ("Почувалися безпечно",f"{o.get('safe_pct',0):g}%"),
            ("Вважають активності корисними",f"{o.get('useful_pct',0):g}%"),
            ("Хочуть прийти ще",f"{o.get('return_pct',0):g}%"),
            ("Середня оцінка",f"{o.get('avg_rating',0):g}/5"),
        ]
        if int(o.get("responses",0))>0:
            fig.text(.07,.835,f"{o.get('high_rating_pct',0):g}% учасників високо оцінили активності.",fontsize=9.5,color=muted)
            fig.text(.07,.812,f"{o.get('new_knowledge_pct',0):g}% повідомили про отримання нових знань; {o.get('safe_pct',0):g}% почувалися безпечно.",fontsize=9.5,color=muted)
            y=.74
            for title,value in impact:
                fig.text(.08,y,wrap_words(title,48),fontsize=11,color=ink,va="center")
                fig.text(.88,y,value,fontsize=15,weight="bold",color=brand,ha="right",va="center")
                y-=.085
            fig.text(.08,.25,f"Завершених анкет зворотного зв’язку: {o.get('responses',0)}",fontsize=10,color=muted)
        else:
            fig.text(.07,.78,"У вибраному періоді завершених feedback-анкет ще немає.",fontsize=14,color=muted)
        fig.text(.07,.15,wrap_words("Показники впливу доповнюють кількісні дані та показують сприйняту користь, навчальний результат, безпеку й намір повернутися.",105),fontsize=9,color=muted)
        footer(fig); pdf.savefig(fig); plt.close(fig)

        # Activity heatmap.
        fig,ax=plt.subplots(figsize=(11.69,8.27)); matrix=data.get("heatmap") or [[0]*24 for _ in range(7)]
        im=ax.imshow(matrix,aspect="auto",interpolation="nearest",cmap="YlGnBu")
        ax.set_yticks(range(7)); ax.set_yticklabels(data.get("heatmap_days",["Пн","Вт","Ср","Чт","Пт","Сб","Нд"]))
        ax.set_xticks(range(24)); ax.set_xticklabels([f"{h:02d}" for h in range(24)],fontsize=8)
        ax.set_xlabel("Година"); ax.set_ylabel("День тижня"); ax.set_title("Теплова карта активності: день тижня × година",fontsize=16,weight="bold",pad=18)
        fig.colorbar(im,ax=ax,fraction=.025,pad=.02,label="Кількість дій")
        fig.tight_layout(); pdf.savefig(fig); plt.close(fig)

        # Reporting 2.0 dynamics: daily for short periods, weekly for medium
        # periods, monthly for long periods. Never draw a meaningless one-point line.
        rows=data.get("trend",data.get("monthly",[]))
        if len(rows)>1:
            fig,ax=plt.subplots(figsize=(11.69,8.27)); x=range(len(rows))
            ax.plot(x,[r["visits"] for r in rows],marker="o",label="Відвідування",color=brand,linewidth=2.2)
            ax.plot(x,[r["new_users"] for r in rows],marker="o",label="Нові учасники",color=accent,linewidth=2.2)
            ax.set_xticks(list(x)); ax.set_xticklabels([r["label"] for r in rows],rotation=35,ha="right"); ax.legend(); ax.grid(axis="y",alpha=.2)
            ax.set_title(f"Динаміка за період {data.get('trend_granularity_label','')}",fontsize=16,weight="bold"); fig.tight_layout(); pdf.savefig(fig); plt.close(fig)
        elif len(rows)==1:
            r=rows[0]
            kpi_page(pdf,[("Відвідування",r["visits"]),("Нові учасники",r["new_users"]),("Завершені події",r["events"]),("Нараховано XP",r["xp"]),("Отримані бейджі",r.get("badges",0))],title="Динаміка за період",subtitle="Є лише одна точка даних — замість лінійного графіка показано KPI.")
        else:
            kpi_page(pdf,[("Відвідування",0),("Нові учасники",0),("Завершені події",0),("Нараховано XP",0)],title="Динаміка за період",subtitle="У вибраному періоді немає даних для побудови динаміки.")

        # Events: all rows, paginated without truncation.
        event_rows=data["events"]
        if not event_rows:
            fig,ax=plt.subplots(figsize=(11.69,8.27)); ax.axis("off"); ax.set_title("Події та відвідуваність",fontsize=16,weight="bold",pad=20); ax.text(.5,.5,"Подій у вибраному періоді немає",ha="center",va="center"); pdf.savefig(fig); plt.close(fig)
        else:
            for page_no,start in enumerate(range(0,len(event_rows),20),1):
                chunk=event_rows[start:start+20]
                fig,ax=plt.subplots(figsize=(11.69,8.27)); ax.axis("off"); ax.set_title(f"Події та відвідуваність — сторінка {page_no}",fontsize=16,weight="bold",pad=20)
                rows=[[r["date"],wrap_words(r["title"],38),r.get("timing_label",r["status"]),str(r["registered"]),str(r["attended"])] for r in chunk]
                table=ax.table(cellText=rows,colLabels=["Дата","Подія","Стан","Зареєстровано","Відвідали"],loc="upper center",cellLoc="left",colWidths=[.17,.43,.14,.13,.13]); table.auto_set_font_size(False); table.set_fontsize(8.2); table.scale(1,1.55)
                for (rr,cc),cell in table.get_celld().items():
                    cell.set_edgecolor("#D4E6E9"); cell.set_linewidth(.55); cell.get_text().set_wrap(True)
                    if rr==0: cell.set_facecolor(brand); cell.set_text_props(color="white",weight="bold",ha="center")
                    elif rr%2==0: cell.set_facecolor("#F7FBFC")
                fig.tight_layout(); pdf.savefig(fig); plt.close(fig)

        # Demographics: paginate every 15 rows so no category disappears.
        for title,key in [("Вікова статистика","age"),("Гендерна статистика","gender"),("Населені пункти","settlement"),("Категорії вразливості — тільки агреговано","vulnerability")]:
            all_items=sorted(data[key].items(),key=lambda x:(-x[1],x[0]))
            chunks=[all_items[i:i+15] for i in range(0,len(all_items),15)] or [[]]
            for page_no,chunk in enumerate(chunks,1):
                fig,ax=plt.subplots(figsize=(11.69,8.27)); _plot_counter(ax,title if len(chunks)==1 else f"{title} — {page_no}/{len(chunks)}",data[key],items=chunk)
                if key=="vulnerability": fig.text(.08,.04,wrap_words(data["privacy_note"],140),fontsize=8)
                fig.tight_layout(rect=[0,.06 if key=="vulnerability" else 0,1,1]); pdf.savefig(fig); plt.close(fig)

        fig,axes=plt.subplots(1,2,figsize=(11.69,8.27)); fig.suptitle("Гейміфікація: ліги та серії",fontsize=16,weight="bold")
        leagues=list(data["league_distribution"].items()); axes[0].barh(range(len(leagues)),[v for _,v in leagues]); axes[0].set_yticks(range(len(leagues))); axes[0].set_yticklabels([k for k,_ in leagues]); axes[0].invert_yaxis(); axes[0].set_title("Ліги — поточний стан"); axes[0].grid(axis="x",alpha=.2)
        ss=data["streak_snapshot"]; streak_items=[("Тижнева",ss["weekly_active"]),("Суперсерія",ss["super_active"]),("Можна відновити",ss["recoverable"])]; axes[1].bar([x[0] for x in streak_items],[x[1] for x in streak_items]); axes[1].set_title("Серії участі"); axes[1].tick_params(axis="x",rotation=20); axes[1].grid(axis="y",alpha=.2); fig.text(.06,.04,f"Заморозка серій у періоді: {data['summary']['streak_freeze_days']} дн.",fontsize=9); fig.tight_layout(rect=[0,.06,1,.93]); pdf.savefig(fig); plt.close(fig)

        fig,ax=plt.subplots(figsize=(11.69,8.27)); rows=data["badge_weekly"]
        if len(rows)>1:
            x=range(len(rows)); ax.plot(x,[r["value"] for r in rows],marker="o"); ax.set_xticks(list(x)); ax.set_xticklabels([r["label"] for r in rows],rotation=35,ha="right"); ax.grid(axis="y",alpha=.2)
        elif len(rows)==1:
            ax.bar([rows[0]["label"]],[rows[0]["value"]]); ax.grid(axis="y",alpha=.2)
        else: ax.text(.5,.5,"У вибраному періоді бейджів не видавали",ha="center",va="center")
        ax.set_title("Отримані бейджі по тижнях",fontsize=16,weight="bold"); fig.tight_layout(); pdf.savefig(fig); plt.close(fig)
    return bio.getvalue()
