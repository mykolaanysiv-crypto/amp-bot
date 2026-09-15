from __future__ import annotations

from .registration_ux import mark_first_activity

import asyncio
import hashlib

from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from secrets import token_urlsafe
from typing import Iterable

import qrcode
from PIL import Image, ImageDraw, ImageFont, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.exc import IntegrityError

from .config import Settings
from .gamification import CLAIMABLE_ACTIVITY_CATALOG, get_level, referral_reward_for_position, normalize_event_xp, normalize_quest_xp
from .profile_data import gender_label, media_consent_label, vulnerability_labels, participant_first_name, split_display_name
from .ui_labels import label, event_registration_status_label
from .security import hash_password
from .runtime_config import ensure_runtime_defaults, get_runtime_int
from .time_utils import event_local_now
from .models import (
    ActivityApplication,
    ActivityType,
    AuditLog,
    BanRecord,
    Badge,
    Event,
    EventRegistration,
    Idea,
    Opportunity,
    RequestCase,
    Quest,
    QuestParticipation,
    Referral,
    Reward,
    RewardClaim,
    Season,
    SurveyResponse,
    ParticipationStreak,
    Team,
    TeamMember,
    User,
    UserBadge,
    UserRole,
    UserStatus,
    VolunteerTask,
    VolunteerTaskParticipation,
    XPTransaction,
    WebStaffAccount,
)


def age_on(birth_date: date, today: date | None = None) -> int:
    today = today or date.today()
    return today.year - birth_date.year - ((today.month, today.day) < (birth_date.month, birth_date.day))


# Startup seeding can be triggered by both the Telegram bot and the web app.
# In run_all.py they share one event loop, so this lock serializes the entire
# bootstrap sequence (not just CREATE TABLE) and prevents duplicate seed rows.
_bootstrap_lock = asyncio.Lock()



async def get_user_by_tg(session: AsyncSession, tg_id: int) -> User | None:
    user = await session.scalar(select(User).where(User.tg_id == tg_id))
    # Restore access immediately when a temporary ban has expired, even before
    # the hourly maintenance job runs. This keeps Telegram behaviour intuitive.
    if user and user.status == UserStatus.BLOCKED.value and user.blocked_until and user.blocked_until <= datetime.utcnow():
        now = datetime.utcnow()
        active_ban = await session.scalar(
            select(BanRecord).where(
                BanRecord.user_id == user.id,
                BanRecord.lifted_at.is_(None),
            ).order_by(BanRecord.started_at.desc())
        )
        if active_ban:
            active_ban.lifted_at = now
            active_ban.lift_reason = "Строк блокування завершився автоматично"
            active_ban.updated_at = now
        user.status = UserStatus.ACTIVE.value
        user.blocked_until = None
        user.block_reason = None
        await session.commit()
    return user


async def get_user(session: AsyncSession, user_id: int) -> User | None:
    return await session.get(User, user_id)


async def ensure_user_tokens(session: AsyncSession, user: User) -> None:
    if not user.public_token:
        user.public_token = token_urlsafe(18)
    if not user.referral_code:
        # compact human-shareable code; uniqueness is additionally enforced by DB index
        user.referral_code = f"AMP{user.id:04d}{token_urlsafe(4).replace('-', '').replace('_', '')[:5]}".upper()


async def current_season(session: AsyncSession) -> Season | None:
    return await session.scalar(select(Season).where(Season.active == True).order_by(Season.starts_at.desc()))  # noqa: E712


async def ensure_default_season(session: AsyncSession, settings: Settings) -> Season:
    """Ensure an initial season without re-activating archived history.

    v1.9.2 makes seasons historical objects. If an administrator has already
    created another active season, startup must respect it instead of forcing
    the config-named season back to active. Likewise, a finalized/expired
    configured season is never resurrected after restart.
    """
    active = await current_season(session)
    if active:
        season = active
    else:
        season = await session.scalar(select(Season).where(Season.name == settings.season_name))
        if not season:
            try:
                async with session.begin_nested():
                    candidate = Season(
                        name=settings.season_name,
                        starts_at=settings.season_start,
                        ends_at=settings.season_end,
                        active=settings.season_end >= date.today(),
                        archived=settings.season_end < date.today(),
                    )
                    session.add(candidate)
                    await session.flush()
            except IntegrityError:
                pass
            season = await session.scalar(select(Season).where(Season.name == settings.season_name))
            if not season:
                raise RuntimeError(f"Не вдалося створити або отримати сезон: {settings.season_name}")
        # Only sync dates for a non-finalized config season. Historical snapshots
        # must remain stable after archival.
        if not season.finalized_at:
            season.starts_at = settings.season_start
            season.ends_at = settings.season_end
            if settings.season_end >= date.today():
                season.active = True
                season.archived = False

    # Backfill legacy XP into whichever season is currently selected.
    start_dt = datetime.combine(season.starts_at, datetime.min.time())
    end_dt = datetime.combine(season.ends_at, datetime.max.time())
    legacy = (await session.scalars(select(XPTransaction).where(
        XPTransaction.season_id.is_(None),
        XPTransaction.created_at >= start_dt,
        XPTransaction.created_at <= end_dt,
    ))).all()
    for tx in legacy:
        tx.season_id = season.id
    return season


async def xp_total(session: AsyncSession, user_id: int) -> int:
    value = await session.scalar(
        select(func.coalesce(func.sum(XPTransaction.amount), 0)).where(XPTransaction.user_id == user_id)
    )
    return int(value or 0)


async def season_xp(session: AsyncSession, user_id: int, season_id: int | None = None) -> int:
    if season_id is None:
        season = await current_season(session)
        if not season:
            return 0
        season_id = season.id
    value = await session.scalar(
        select(func.coalesce(func.sum(XPTransaction.amount), 0)).where(
            XPTransaction.user_id == user_id,
            XPTransaction.season_id == season_id,
        )
    )
    return int(value or 0)


async def add_xp(
    session: AsyncSession,
    user: User,
    amount: int,
    description: str,
    category: str = "other",
    created_by: int | None = None,
    event_id: int | None = None,
) -> tuple[int, str, bool]:
    before = await xp_total(session, user.id)
    before_level = get_level(before)[0]
    season = await current_season(session)
    tx = XPTransaction(
        user_id=user.id,
        amount=amount,
        description=description,
        category=category,
        created_by=created_by,
        event_id=event_id,
        season_id=season.id if season else None,
    )
    session.add(tx)
    # Spendable wallet follows earned/corrected XP, but reward redemption does not
    # change lifetime XP or level.
    user.wallet_xp = max(0, int(user.wallet_xp or 0) + amount)
    await session.flush()
    if amount >= 0 and category in {"event", "quest", "task", "activity", "survey", "team_quest", "idea_approved"}:
        await mark_first_activity(session, user.id, tx.created_at or datetime.utcnow())
    after = before + amount
    after_level = get_level(after)[0]
    await evaluate_automatic_badges(session, user)
    return after, after_level, before_level != after_level


async def process_birthdays(session: AsyncSession, today: date) -> list[tuple[int, str, int]]:
    """Award the annual birthday bonus exactly once per calendar year.

    Returns tuples: (telegram_id, first_name, xp_awarded). Only active users with
    a birth date are eligible. The persisted birthday_reward_year makes the job
    safe across restarts and repeated scheduler runs.
    """
    birthday_xp = await get_runtime_int(session, "xp.birthday")
    users = (await session.scalars(
        select(User).where(
            User.status == UserStatus.ACTIVE.value,
            User.birth_date.is_not(None),
        )
    )).all()
    rewarded: list[tuple[int, str, int]] = []
    for user in users:
        if not user.birth_date:
            continue
        if (user.birth_date.month, user.birth_date.day) != (today.month, today.day):
            continue
        if user.birthday_reward_year == today.year:
            continue
        await add_xp(
            session,
            user,
            birthday_xp,
            "Подарунок АМП до дня народження 🎂",
            category="birthday",
        )
        user.birthday_reward_year = today.year
        first_name = participant_first_name(user)
        rewarded.append((user.tg_id, first_name, birthday_xp))
    return rewarded


async def process_expired_bans(session: AsyncSession, now: datetime | None = None) -> int:
    """Automatically reactivate users whose temporary ban has expired."""
    now = now or datetime.utcnow()
    users = (await session.scalars(
        select(User).where(
            User.status == UserStatus.BLOCKED.value,
            User.blocked_until.is_not(None),
            User.blocked_until <= now,
        )
    )).all()
    for user in users:
        active_ban = await session.scalar(
            select(BanRecord).where(
                BanRecord.user_id == user.id,
                BanRecord.lifted_at.is_(None),
            ).order_by(BanRecord.started_at.desc())
        )
        if active_ban:
            active_ban.lifted_at = now
            active_ban.lift_reason = "Строк блокування завершився автоматично"
            active_ban.updated_at = now
        user.status = UserStatus.ACTIVE.value
        user.blocked_until = None
        user.block_reason = None
        await log_audit(session, "temporary_ban_expired", actor_label="system", entity_type="user", entity_id=user.id, details="Тимчасове блокування завершено автоматично")
    return len(users)


async def create_event(
    session: AsyncSession,
    title: str,
    description: str,
    starts_at: datetime,
    location: str,
    xp_reward: int,
    volunteer_hours: float,
    created_by: int,
) -> Event:
    xp_reward = normalize_event_xp(xp_reward)
    event = Event(
        title=title,
        description=description,
        starts_at=starts_at,
        location=location,
        xp_reward=xp_reward,
        volunteer_hours=volunteer_hours,
        checkin_token=token_urlsafe(18),
        share_token=token_urlsafe(18),
        created_by=created_by,
    )
    session.add(event)
    await session.flush()
    return event


async def register_for_event(session: AsyncSession, user_id: int, event_id: int) -> EventRegistration:
    reg = await session.scalar(
        select(EventRegistration).where(EventRegistration.user_id == user_id, EventRegistration.event_id == event_id)
    )
    if reg:
        reg.status = "registered"
        reg.registered_at = datetime.utcnow()
        reg.waitlisted_at = None
        reg.reservation_expires_at = None
        reg.no_show_at = None
        return reg
    reg = EventRegistration(user_id=user_id, event_id=event_id, status="registered")
    session.add(reg)
    await session.flush()
    return reg


EVENT_OCCUPIED_STATUSES = {"registered", "reserved", "checked_in", "attended"}


async def join_event_waitlist(session: AsyncSession, user_id: int, event_id: int, *, now: datetime | None = None) -> EventRegistration:
    """Place a participant in an event queue without consuming event capacity."""
    now = now or datetime.utcnow()
    reg = await session.scalar(
        select(EventRegistration).where(EventRegistration.user_id == user_id, EventRegistration.event_id == event_id)
    )
    if not reg:
        reg = EventRegistration(user_id=user_id, event_id=event_id, status="waitlisted", waitlisted_at=now)
        session.add(reg)
        await session.flush()
        return reg
    if reg.status == "attended":
        return reg
    reg.status = "waitlisted"
    reg.waitlisted_at = now
    reg.waitlist_promoted_at = None
    reg.reservation_expires_at = None
    reg.checkin_at = None
    reg.no_show_at = None
    return reg


async def accept_event_reservation(session: AsyncSession, user_id: int, event_id: int, *, now: datetime | None = None) -> tuple[EventRegistration | None, str]:
    """Confirm a two-hour waitlist reservation. Returns (registration, state)."""
    now = now or datetime.utcnow()
    reg = await session.scalar(
        select(EventRegistration).where(EventRegistration.user_id == user_id, EventRegistration.event_id == event_id)
    )
    if not reg or reg.status != "reserved":
        return reg, "not_reserved"
    if reg.reservation_expires_at and reg.reservation_expires_at <= now:
        reg.status = "waitlisted"
        reg.waitlisted_at = now
        reg.reservation_expires_at = None
        return reg, "expired"
    reg.status = "registered"
    reg.registered_at = now
    reg.reservation_expires_at = None
    return reg, "accepted"


async def process_event_operations(session: AsyncSession, *, now: datetime | None = None, event_id: int | None = None) -> dict[str, int]:
    """Maintain waitlist reservations and attendance terminal states.

    - Expired two-hour reservations return to the end of the queue.
    - Free capacity promotes the oldest queued participant and reserves a place for two hours.
    - Once an event is completed, participants who remained merely registered/reserved become no-shows.

    Notification Center provides durable Telegram delivery and deduplication.
    """
    from .reliability import queue_telegram_delivery

    explicit_now = now
    now = now or datetime.utcnow()
    event_now = explicit_now or event_local_now()
    changed = {"expired_reservations": 0, "promoted_waitlist": 0, "cancelled_waitlist": 0, "no_show": 0}

    expired_stmt = select(EventRegistration).where(
        EventRegistration.status == "reserved",
        EventRegistration.reservation_expires_at.is_not(None),
        EventRegistration.reservation_expires_at <= now,
    )
    if event_id is not None:
        expired_stmt = expired_stmt.where(EventRegistration.event_id == int(event_id))
    expired = list((await session.scalars(expired_stmt)).all())
    for reg in expired:
        reg.status = "waitlisted"
        reg.waitlisted_at = now
        reg.reservation_expires_at = None
        changed["expired_reservations"] += 1
        user = await session.get(User, reg.user_id)
        event = await session.get(Event, reg.event_id)
        if user and event and user.status == UserStatus.ACTIVE.value and user.tg_id:
            await queue_telegram_delivery(
                session, user.tg_id,
                f"⏳ <b>Резерв місця завершився</b>\n\nДвогодинний резерв на подію «<b>{event.title}</b>» минув. "
                "Ми повернули тебе в чергу. Якщо звільниться наступне місце — бот повідомить автоматично.",
                source="event_waitlist",
                dedupe_key=f"event_waitlist_expired:{event.id}:{reg.id}:{int(now.timestamp())//7200}",
            )

    events_stmt = select(Event).where(
        Event.capacity.is_not(None), Event.capacity > 0, Event.status.in_(["open", "closed", "postponed"])
    )
    if event_id is not None:
        events_stmt = events_stmt.where(Event.id == int(event_id))
    events = list((await session.scalars(events_stmt.order_by(Event.starts_at.asc()))).all())
    checkin_close_minutes = await get_runtime_int(session, "events.checkin_close_after_minutes")
    for event in events:
        # Do not promote people after the configured operational window has ended.
        if event.cancelled_at or (event.starts_at and event.starts_at + timedelta(minutes=checkin_close_minutes) < event_now):
            continue
        occupied = int(await session.scalar(
            select(func.count(EventRegistration.id)).where(
                EventRegistration.event_id == event.id,
                EventRegistration.status.in_(EVENT_OCCUPIED_STATUSES),
            )
        ) or 0)
        while occupied < int(event.capacity or 0):
            reg = await session.scalar(
                select(EventRegistration).where(
                    EventRegistration.event_id == event.id,
                    EventRegistration.status == "waitlisted",
                ).order_by(EventRegistration.waitlisted_at.asc().nullsfirst(), EventRegistration.id.asc())
            )
            if not reg:
                break
            user = await session.get(User, reg.user_id)
            if not user or user.status != UserStatus.ACTIVE.value:
                reg.status = "cancelled"
                changed["cancelled_waitlist"] += 1
                continue
            reg.status = "reserved"
            reg.waitlist_promoted_at = now
            reservation_minutes = await get_runtime_int(session, "events.waitlist_reservation_minutes")
            reg.reservation_expires_at = now + timedelta(minutes=reservation_minutes)
            changed["promoted_waitlist"] += 1
            occupied += 1
            await queue_telegram_delivery(
                session, user.tg_id,
                f"🎉 <b>Звільнилося місце!</b>\n\nНа подію «<b>{event.title}</b>» для тебе зарезервовано місце на <b>{reservation_minutes} хв</b>. "
                "Підтвердь його кнопкою нижче, інакше резерв перейде наступному учаснику в черзі.",
                source="event_waitlist",
                dedupe_key=f"event_waitlist_reserved:{event.id}:{reg.id}:{int(now.timestamp())//7200}",
                button_text="✅ Підтвердити місце",
                callback_data=f"event_reserve_accept:{event.id}",
            )

    completed_stmt = select(Event.id).where(Event.status == "completed")
    if event_id is not None:
        completed_stmt = completed_stmt.where(Event.id == int(event_id))
    completed_events = list((await session.scalars(completed_stmt)).all())
    if completed_events:
        regs = list((await session.scalars(
            select(EventRegistration).where(
                EventRegistration.event_id.in_(completed_events),
                EventRegistration.status.in_(["registered", "reserved"]),
            )
        )).all())
        for reg in regs:
            reg.status = "no_show"
            reg.no_show_at = now
            reg.reservation_expires_at = None
            changed["no_show"] += 1

    return changed


async def event_checkin_window(
    session: AsyncSession, event: Event, *, now: datetime | None = None
) -> dict[str, object]:
    """Return the operational check-in/attendance window for an event.

    Datetimes in the current schema are naive and are compared consistently with
    the rest of the application. Runtime settings make the window adjustable
    without a deploy.
    """
    now = now or event_local_now()
    before = await get_runtime_int(session, "events.checkin_open_before_minutes")
    after = await get_runtime_int(session, "events.checkin_close_after_minutes")
    opens_at = event.starts_at - timedelta(minutes=before)
    closes_at = event.starts_at + timedelta(minutes=after)
    if now < opens_at:
        state = "too_early"
    elif now > closes_at:
        state = "closed"
    else:
        state = "open"
    return {"state": state, "opens_at": opens_at, "closes_at": closes_at, "now": now, "before_minutes": before, "after_minutes": after}


async def checkin_for_event(
    session: AsyncSession, user_id: int, token: str, *, now: datetime | None = None
) -> tuple[Event | None, str]:
    event = await session.scalar(select(Event).where(Event.checkin_token == token))
    if not event or event.status not in {"open", "closed", "postponed"}:
        return None, "invalid"
    access = await event_checkin_window(session, event, now=now)
    if access["state"] == "too_early":
        return event, "too_early"
    if access["state"] == "closed":
        return event, "window_closed"
    reg = await session.scalar(
        select(EventRegistration).where(EventRegistration.user_id == user_id, EventRegistration.event_id == event.id)
    )
    if not reg:
        reg = EventRegistration(user_id=user_id, event_id=event.id, status="checked_in")
        session.add(reg)
    elif reg.status != "attended":
        reg.status = "checked_in"
    reg.checkin_at = access["now"]
    return event, "ok"


async def confirm_single_event_attendance(
    session: AsyncSession, event: Event, reg: EventRegistration, admin_user: User,
    *, override_reason: str | None = None, now: datetime | None = None,
) -> tuple[User, int, str, bool] | None:
    """Confirm attendance and award benefits exactly once.

    Ordinary confirmation is allowed only inside the configured attendance
    window. Staff may bypass it only by supplying a non-empty override reason;
    callers are responsible for writing that reason to the audit log.
    """
    if reg.event_id != event.id or reg.status == "attended":
        return None
    if reg.status != "checked_in":
        return None
    access = await event_checkin_window(session, event, now=now)
    override = (override_reason or "").strip()
    if access["state"] != "open" and len(override) < 5:
        return None

    # Lock/reload the registration on databases that support row locking. This
    # closes the common concurrent double-confirm path before XP is inserted.
    locked = await session.scalar(
        select(EventRegistration)
        .where(EventRegistration.id == reg.id)
        .with_for_update()
        .execution_options(populate_existing=True)
    )
    if not locked or locked.status == "attended" or locked.status != "checked_in":
        return None
    reg = locked
    user = await session.get(User, reg.user_id)
    if not user:
        return None
    event.xp_reward = normalize_event_xp(event.xp_reward)
    total, level, leveled = await add_xp(
        session,
        user,
        event.xp_reward,
        f"Участь у події «{event.title}»",
        category="event",
        created_by=admin_user.id,
        event_id=event.id,
    )
    user.volunteer_hours += event.volunteer_hours
    reg.status = "attended"
    reg.confirmed_at = access["now"]
    reg.attendance_confirmed_by_user_id = admin_user.id
    if not reg.attendance_signature:
        raw_signature = "|".join([
            "AMP-ATTENDANCE-v1", str(event.id), str(reg.id), str(user.id),
            (reg.checkin_at or reg.confirmed_at).isoformat(), reg.confirmed_at.isoformat(), token_urlsafe(24),
        ])
        reg.attendance_signature = hashlib.sha256(raw_signature.encode("utf-8")).hexdigest()
        reg.attendance_signature_version = "sha256-v1"
        reg.attendance_signature_created_at = reg.confirmed_at
    await evaluate_automatic_badges(session, user)
    return user, total, level, leveled


async def admin_scan_event_participant(
    session: AsyncSession,
    event_id: int,
    participant_id: int,
    admin_user: User,
    *,
    allow_register: bool = False,
) -> dict:
    """Resolve a participant QR/AMP-ID and confirm attendance from an admin scanner.

    The workflow is deliberately shared by Telegram scanner flows and can be used
    by web integrations. It preserves the existing idempotent attendance award.
    """
    event = await session.get(Event, event_id)
    if not event or event.cancelled_at or event.status in {"draft", "cancelled", "completed"}:
        return {"ok": False, "code": "event_unavailable", "message": "Відмітка для цієї події недоступна."}

    access = await event_checkin_window(session, event)
    if access["state"] != "open":
        return {
            "ok": False, "code": str(access["state"]),
            "message": "Відмітку ще не відкрито." if access["state"] == "too_early" else "Вікно відмітки та підтвердження участі вже закрито. Для винятку використайте ручне підтвердження з обов’язковою причиною.",
            "event": event, "window": access,
        }

    user = await session.get(User, participant_id)
    if not user:
        return {"ok": False, "code": "user_not_found", "message": "Учасника не знайдено."}
    if user.status != UserStatus.ACTIVE.value:
        return {
            "ok": False, "code": "user_inactive", "message": f"Акаунт учасника не активний: {label(user.status)}.",
            "event": event, "user": user,
        }

    reg = await session.scalar(select(EventRegistration).where(
        EventRegistration.event_id == event.id, EventRegistration.user_id == user.id
    ))
    if reg and reg.status == "attended":
        return {
            "ok": True, "code": "already_attended", "message": "Присутність уже була підтверджена раніше.",
            "event": event, "user": user, "registration": reg,
        }

    eligible = bool(reg and reg.status in {"registered", "reserved", "checked_in"})
    if not eligible and not allow_register:
        return {
            "ok": True, "code": "unregistered", "message": "Учасник не зареєстрований на цю подію.",
            "event": event, "user": user, "registration": reg, "requires_registration": True,
        }

    now = event_local_now()
    if not reg:
        reg = EventRegistration(event_id=event.id, user_id=user.id, status="registered", registered_at=now)
        session.add(reg)
        await session.flush()
    elif allow_register and reg.status not in {"registered", "reserved", "checked_in"}:
        reg.status = "registered"
        reg.registered_at = now
        reg.waitlisted_at = None
        reg.reservation_expires_at = None
        reg.no_show_at = None

    if reg.status == "reserved":
        reg.registered_at = now
        reg.reservation_expires_at = None
    if reg.status != "checked_in":
        reg.status = "checked_in"
        reg.checkin_at = now

    result = await confirm_single_event_attendance(session, event, reg, admin_user)
    if not result:
        return {"ok": False, "code": "confirm_failed", "message": "Не вдалося підтвердити присутність.", "event": event, "user": user, "registration": reg}
    confirmed_user, total, level_name, leveled = result
    return {
        "ok": True, "code": "confirmed", "message": "Присутність підтверджено",
        "event": event, "user": confirmed_user, "registration": reg,
        "total_xp": total, "level_name": level_name, "leveled": leveled,
        "xp": event.xp_reward, "hours": event.volunteer_hours,
    }


async def confirm_event_attendance(session: AsyncSession, event: Event, admin_user: User, *, override_reason: str | None = None, now: datetime | None = None) -> tuple[int, list[tuple[User, int, str, bool]]]:
    regs = (
        await session.scalars(
            select(EventRegistration).where(EventRegistration.event_id == event.id, EventRegistration.status == "checked_in")
        )
    ).all()
    results: list[tuple[User, int, str, bool]] = []
    for reg in regs:
        result = await confirm_single_event_attendance(session, event, reg, admin_user, override_reason=override_reason, now=now)
        if result:
            results.append(result)
    return len(results), results


async def seed_activity_types(session: AsyncSession) -> None:
    """Seed missing balanced activity types without overwriting admin edits."""
    for order, item in enumerate(CLAIMABLE_ACTIVITY_CATALOG, start=10):
        row = await session.scalar(select(ActivityType).where(ActivityType.code == item["code"]))
        if row:
            continue
        try:
            async with session.begin_nested():
                session.add(ActivityType(
                    code=item["code"], title=item["title"], category=item["category"],
                    description=item["description"], instructions=item["instructions"],
                    xp_reward=int(item["xp"]), hours_reward=float(item["hours"]),
                    active=True, sort_order=order,
                ))
                await session.flush()
        except IntegrityError:
            pass


async def complete_activity_application(
    session: AsyncSession, application: ActivityApplication, admin_user: User | None = None
) -> tuple[User, int, str, bool] | None:
    """Complete an approved/submitted application and award its snapshotted XP exactly once."""
    if application.status == "activity_completed":
        return None
    if application.status not in {"activity_approved", "activity_submitted"}:
        return None
    user = await session.get(User, application.user_id)
    activity = await session.get(ActivityType, application.activity_type_id)
    if not user or not activity:
        return None
    result = await add_xp(
        session, user, int(application.xp_reward or activity.xp_reward),
        f"Активність «{activity.title}»", category="activity",
        created_by=admin_user.id if admin_user else None,
    )
    user.volunteer_hours += float(application.hours_reward or activity.hours_reward or 0)
    application.status = "activity_completed"
    application.completed_at = datetime.utcnow()
    application.completed_by = admin_user.id if admin_user else None
    await evaluate_automatic_badges(session, user)
    return (user, *result)


async def seed_badges(session: AsyncSession) -> None:
    defaults = [
        ("Перший крок", "🚀", "Перша підтверджена активність", "xp_transactions", 1),
        ("Прокачаний", "🎓", "10 підтверджених активностей", "xp_transactions", 10),
        ("Серце команди", "❤️", "50 волонтерських годин", "volunteer_hours", 50),
        ("100 годин для АМП", "⏱", "100 волонтерських годин", "volunteer_hours", 100),
        ("Магніт", "👥", "3 успішно залучені нові учасники", "referrals", 3),
        ("Квестер", "🎯", "5 підтверджених квестів", "quests", 5),
        ("АМПасадор", "🔥", "Досягнення 300 XP", "xp_total", 300),
        ("Лідер АМП", "🛰️", "Досягнення 800 XP", "xp_total", 800),
        ("Легенда АМП", "🏆", "Досягнення 1200 XP", "xp_total", 1200),
    ]
    for name, icon, desc, criteria_type, criteria_value in defaults:
        badge = await session.scalar(select(Badge).where(Badge.name == name))
        if not badge:
            try:
                async with session.begin_nested():
                    session.add(Badge(
                        name=name,
                        icon=icon,
                        description=desc,
                        criteria_type=criteria_type,
                        criteria_value=criteria_value,
                        automatic=True,
                    ))
                    await session.flush()
            except IntegrityError:
                pass
            badge = await session.scalar(select(Badge).where(Badge.name == name))
        if badge:
            badge.icon = icon
            badge.description = desc
            badge.criteria_type = criteria_type
            badge.criteria_value = criteria_value
            badge.automatic = True

    # Manual / thematic badges remain available to admins.
    manual = [
        ("Чистий старт", "🧹", "Участь у толоках та благоустрої"),
        ("Голос АМП", "🎤", "Проведення власної активності"),
        ("Контент-мейкер", "📸", "Внесок у комунікації та медіа"),
        ("Нетворкер", "🤝", "Залучення партнерів"),
        ("Ідейник", "💡", "Реалізовані ідеї"),
        ("Ментор", "🧑‍🏫", "Допомога новим учасникам"),
        ("Запускаю зміни", "🚀", "Реалізація власного мініпроєкту"),
    ]
    for name, icon, desc in manual:
        badge = await session.scalar(select(Badge).where(Badge.name == name))
        if not badge:
            try:
                async with session.begin_nested():
                    session.add(Badge(name=name, icon=icon, description=desc, automatic=False))
                    await session.flush()
            except IntegrityError:
                pass


async def _metric_value(session: AsyncSession, user: User, criteria_type: str) -> int:
    if criteria_type == "xp_total":
        return await xp_total(session, user.id)
    if criteria_type == "volunteer_hours":
        return int(user.volunteer_hours or 0)
    if criteria_type == "xp_transactions":
        return int(await session.scalar(select(func.count(XPTransaction.id)).where(XPTransaction.user_id == user.id)) or 0)
    if criteria_type == "referrals":
        return int(await session.scalar(select(func.count(Referral.id)).where(Referral.inviter_user_id == user.id, Referral.status == "rewarded")) or 0)
    if criteria_type == "quests":
        return int(await session.scalar(select(func.count(QuestParticipation.id)).where(QuestParticipation.user_id == user.id, QuestParticipation.status == "approved")) or 0)
    if criteria_type == "events":
        return int(await session.scalar(select(func.count(EventRegistration.id)).where(EventRegistration.user_id == user.id, EventRegistration.status == "attended")) or 0)
    if criteria_type == "activities":
        return int(await session.scalar(select(func.count(ActivityApplication.id)).where(ActivityApplication.user_id == user.id, ActivityApplication.status == "activity_completed")) or 0)
    if criteria_type == "ideas":
        return int(await session.scalar(select(func.count(Idea.id)).where(Idea.user_id == user.id, Idea.status == "implemented")) or 0)
    if criteria_type == "tasks":
        return int(await session.scalar(select(func.count(VolunteerTaskParticipation.id)).where(VolunteerTaskParticipation.user_id == user.id, VolunteerTaskParticipation.status == "approved")) or 0)
    if criteria_type == "surveys":
        return int(await session.scalar(select(func.count(SurveyResponse.id)).where(SurveyResponse.user_id == user.id)) or 0)
    if criteria_type in {"weekly_streak", "event_streak"}:
        streak = await session.scalar(select(ParticipationStreak).where(ParticipationStreak.user_id == user.id))
        return int(getattr(streak, criteria_type, 0) or 0) if streak else 0
    return 0


async def evaluate_automatic_badges(session: AsyncSession, user: User) -> list[Badge]:
    badges = (await session.scalars(select(Badge).where(Badge.active == True, Badge.automatic == True))).all()  # noqa: E712
    awarded: list[Badge] = []
    for badge in badges:
        if badge.badge_type == "ambassador" and user.role not in {"ambassador","coordinator","admin","superadmin"}:
            continue
        if not badge.criteria_type or badge.criteria_value is None:
            continue
        exists = await session.scalar(select(UserBadge).where(UserBadge.user_id == user.id, UserBadge.badge_id == badge.id))
        if exists:
            continue
        if badge.criteria_type in {"donation_first", "donation_single", "donation_total_over"}:
            # Donation badges have slightly different semantics from ordinary >= metrics:
            # first/single use the largest qualifying donation, while cumulative badges
            # are intentionally strict "more than" thresholds.
            from .donations import donation_totals_for_user
            donation_total, donation_largest, donation_count = await donation_totals_for_user(session, user.id)
            if badge.criteria_type in {"donation_first", "donation_single"}:
                qualifies = donation_count > 0 and donation_largest >= int(badge.criteria_value)
            else:
                qualifies = donation_total > int(badge.criteria_value)
        else:
            value = await _metric_value(session, user, badge.criteria_type)
            qualifies = value >= badge.criteria_value
        if qualifies:
            session.add(UserBadge(user_id=user.id, badge_id=badge.id, awarded_by=None))
            awarded.append(badge)
    if awarded:
        await session.flush()
    return awarded


async def ensure_superadmins(session: AsyncSession, ids: set[int]) -> None:
    """Ensure configured Telegram IDs exist and have superadmin rights.

    The function is idempotent. The nested transaction also makes the insert
    tolerant of a concurrent bootstrap in another process (for example bot +
    web containers starting against the same PostgreSQL database).
    """
    for tg_id in ids:
        user = await get_user_by_tg(session, tg_id)
        if not user:
            try:
                async with session.begin_nested():
                    candidate = User(
                        tg_id=tg_id,
                        username=None,
                        full_name=f"Суперадміністратор {tg_id}",
                        role=UserRole.SUPERADMIN.value,
                        status=UserStatus.ACTIVE.value,
                    )
                    session.add(candidate)
                    await session.flush()
            except IntegrityError:
                # Another startup worker inserted the same tg_id first.
                pass
            user = await get_user_by_tg(session, tg_id)

        if user:
            user.role = UserRole.SUPERADMIN.value
            user.status = UserStatus.ACTIVE.value
            await ensure_user_tokens(session, user)




async def seed_streak_restore_reward(session: AsyncSession) -> None:
    existing = await session.scalar(select(Reward).where(Reward.reward_type == "streak_restore"))
    cost = await get_runtime_int(session, "xp.streak_restore_cost")
    if existing:
        existing.min_xp = cost
        return
    session.add(Reward(
        title="Повернути суперсерію",
        description="Відновлює останню втрачену суперсерію відвідування подій. Працює лише якщо є серія, доступна для відновлення.",
        min_xp=cost,
        stock=None,
        active=True,
        reward_type="streak_restore",
    ))

async def ensure_web_staff_accounts(session: AsyncSession, settings: Settings) -> None:
    """Migrate legacy env credentials into hashed DB accounts once.

    Existing DB accounts are never overwritten from environment variables. This
    makes WEB_ADMIN_PASSWORD / WEB_STAFF_ACCOUNTS_JSON bootstrap-only and allows
    operators to remove them from Heroku after the first successful v1.7.3 boot.
    """
    super_username = (settings.web_admin_username or "admin").strip()
    if super_username:
        row = await session.scalar(select(WebStaffAccount).where(WebStaffAccount.username == super_username))
        if not row and settings.web_admin_password:
            tg_id = min(settings.superadmin_ids) if settings.superadmin_ids else None
            session.add(WebStaffAccount(
                username=super_username,
                password_hash=hash_password(settings.web_admin_password),
                display_name="Суперадміністратор",
                role="superadmin",
                active=True,
                must_change_password=True,
                two_factor_enabled=True,
                two_factor_tg_id=tg_id,
                password_changed_at=datetime.utcnow(),
                updated_at=datetime.utcnow(),
            ))

    for username, cfg in (settings.web_staff_accounts or {}).items():
        username = (username or "").strip()
        if not username:
            continue
        exists = await session.scalar(select(WebStaffAccount).where(WebStaffAccount.username == username))
        if exists:
            continue
        password = str(cfg.get("password") or "")
        if not password:
            continue
        session.add(WebStaffAccount(
            username=username,
            password_hash=hash_password(password),
            display_name=str(cfg.get("display_name") or username),
            role="admin",
            active=True,
            must_change_password=True,
            two_factor_enabled=False,
            password_changed_at=datetime.utcnow(),
            updated_at=datetime.utcnow(),
        ))


DEFAULT_SPACE_REWARDS = (
    ("1 год оренди кімнати в молодіжному просторі", "Оренда окремої кімнати АМП на 1 годину за попереднім погодженням з командою простору.", 15),
    ("1 год оренди всього молодіжного простору", "Оренда всього молодіжного простору на 1 годину за попереднім погодженням.", 50),
    ("4 год оренди молодіжного центру", "Оренда молодіжного центру на 4 години за попереднім погодженням.", 150),
    ("Оренда проєктора та екрану", "Оренда проєктора та екрану за попереднім погодженням і правилами користування обладнанням.", 45),
    ("Настільні ігри додому на 7 днів", "Можна взяти доступний набір настільних ігор додому на строк до 7 днів.", 100),
    ("Нова гра на PlayStation рівня AA або старше 10 років", "Запит на придбання нової гри рівня AA або гри, старшої за 10 років. Остаточний вибір погоджує команда АМП.", 50),
    ("Нова гра на PlayStation рівня AAA", "Запит на придбання нової гри рівня AAA. Остаточний вибір погоджує команда АМП.", 75),
    ("Гра на PlayStation — ексклюзив або новинка до 1 року", "Запит на придбання ексклюзиву або новинки до 1 року. Остаточний вибір погоджує команда АМП.", 100),
    ("1 година гри на приставці", "1 година гри на приставці в молодіжному просторі за правилами АМП.", 5),
)


async def seed_default_space_rewards(session: AsyncSession) -> None:
    """Add the v1.7.4.1 default reward catalog without overwriting admin edits."""
    for title, description, cost in DEFAULT_SPACE_REWARDS:
        existing = await session.scalar(select(Reward).where(Reward.title == title))
        if existing:
            continue
        session.add(Reward(title=title, description=description, min_xp=cost, stock=None, active=True, reward_type="service"))


async def ensure_event_share_tokens(session: AsyncSession) -> int:
    """Backfill safe public-share tokens for legacy events without exposing check-in tokens."""
    rows = (await session.scalars(select(Event).where(Event.share_token.is_(None)))).all()
    for event in rows:
        event.share_token = token_urlsafe(18)
    if rows:
        await session.flush()
    return len(rows)


async def bootstrap_defaults(db, settings: Settings) -> None:
    """Create/update startup defaults exactly once per process.

    Both the bot and FastAPI app call this helper. The lock ensures the second
    caller waits until the first transaction has committed, so it sees the
    already-created superadmin, season, badges and default team.
    """
    async with _bootstrap_lock:
        async with db.session_factory() as session:
            from .settlements import ensure_settlement_directory
            from .donations import ensure_donation_badges
            await ensure_runtime_defaults(session)
            await ensure_settlement_directory(session)
            await ensure_superadmins(session, settings.superadmin_ids)
            await ensure_web_staff_accounts(session, settings)
            await ensure_default_season(session, settings)
            await seed_badges(session)
            await ensure_donation_badges(session)
            await seed_activity_types(session)
            await seed_streak_restore_reward(session)
            await seed_default_space_rewards(session)
            await ensure_event_share_tokens(session)
            await add_active_users_to_default_team(session)
            await session.commit()


async def create_referral_for_user(session: AsyncSession, new_user: User, referral_code: str | None) -> Referral | None:
    if not referral_code:
        return None
    inviter = await session.scalar(select(User).where(User.referral_code == referral_code))
    if not inviter or inviter.id == new_user.id:
        return None
    new_user.referred_by_user_id = inviter.id
    referral = Referral(inviter_user_id=inviter.id, invited_user_id=new_user.id, status="pending")
    session.add(referral)
    await session.flush()
    return referral


def referral_quarter_bounds(now: datetime | None = None) -> tuple[datetime, datetime]:
    """Return [quarter_start, next_quarter_start) for the supplied UTC datetime."""
    now = now or datetime.utcnow()
    start_month = ((now.month - 1) // 3) * 3 + 1
    start = datetime(now.year, start_month, 1)
    if start_month == 10:
        end = datetime(now.year + 1, 1, 1)
    else:
        end = datetime(now.year, start_month + 3, 1)
    return start, end


async def referral_quarter_summary(
    session: AsyncSession, inviter_user_id: int, now: datetime | None = None
) -> tuple[int, int, datetime, datetime]:
    """Return rewarded referrals this quarter and XP for the next successful one."""
    start, end = referral_quarter_bounds(now)
    count = int(await session.scalar(
        select(func.count(Referral.id)).where(
            Referral.inviter_user_id == inviter_user_id,
            Referral.status == "rewarded",
            Referral.rewarded_at.is_not(None),
            Referral.rewarded_at >= start,
            Referral.rewarded_at < end,
        )
    ) or 0)
    base_reward = await get_runtime_int(session, "xp.referral_max")
    reward = max(1, int(base_reward) + 1 - max(1, count + 1))
    return count, reward, start, end


async def reward_referral_if_ready(
    session: AsyncSession, invited_user: User, settings: Settings, created_by: int | None = None
) -> tuple[User, int] | None:
    """Reward a successful referral using the quarterly diminishing schedule.

    1st referral in a calendar quarter = 10 XP, 2nd = 9 ... 10th = 1,
    and every subsequent successful referral in that quarter = 1 XP.
    The counter resets automatically on Jan 1, Apr 1, Jul 1 and Oct 1.
    """
    referral = await session.scalar(select(Referral).where(Referral.invited_user_id == invited_user.id))
    if not referral or referral.status != "pending" or invited_user.status != UserStatus.ACTIVE.value:
        return None
    inviter = await session.get(User, referral.inviter_user_id)
    if not inviter:
        return None
    count, reward, _, _ = await referral_quarter_summary(session, inviter.id)
    referral.xp_reward = reward
    referral.status = "rewarded"
    referral.rewarded_at = datetime.utcnow()
    await add_xp(
        session,
        inviter,
        reward,
        f"Запрошено нового активного учасника: {invited_user.full_name} (№{count + 1} у кварталі)",
        category="referral",
        created_by=created_by,
    )
    await evaluate_automatic_badges(session, inviter)
    return inviter, reward


async def revoke_referral_reward_if_inactive(
    session: AsyncSession,
    invited_user: User,
    *,
    reason: str = "Учасник став неактивним",
    now: datetime | None = None,
    window_days: int = 30,
) -> tuple[User, int, int] | None:
    """Claw back referral XP when the invited account becomes inactive within 30 days.

    The operation is idempotent: only referrals in ``rewarded`` state can be revoked.
    Returns (inviter, xp_removed, days_after_reward) when a clawback was applied.
    """
    now = now or datetime.utcnow()
    referral = await session.scalar(select(Referral).where(Referral.invited_user_id == invited_user.id))
    if not referral or referral.status != "rewarded" or not referral.rewarded_at or int(referral.xp_reward or 0) <= 0:
        return None
    elapsed = now - referral.rewarded_at
    if elapsed.total_seconds() < 0 or elapsed > timedelta(days=max(1, int(window_days))):
        return None
    inviter = await session.get(User, referral.inviter_user_id)
    if not inviter:
        return None
    amount = int(referral.xp_reward or 0)
    days_after = max(0, elapsed.days)
    await add_xp(
        session, inviter, -amount,
        f"Повернення бонусу за запрошення: {invited_user.full_name} став(ла) неактивним(ою) протягом {window_days} днів",
        category="referral_reversal",
    )
    referral.status = "revoked"
    referral.revoked_at = now
    referral.revoke_reason = (reason or "Учасник став неактивним")[:500]
    referral.clawback_xp = amount
    await log_audit(
        session, "referral_reward_revoked", actor_label="system", entity_type="referral", entity_id=referral.id,
        details=f"invited_user={invited_user.id}; inviter={inviter.id}; xp=-{amount}; days={days_after}; reason={referral.revoke_reason}",
    )
    return inviter, amount, days_after


async def log_audit(
    session: AsyncSession,
    action: str,
    actor: User | None = None,
    actor_label: str | None = None,
    entity_type: str | None = None,
    entity_id: int | None = None,
    details: str = "",
) -> AuditLog:
    row = AuditLog(
        actor_user_id=actor.id if actor else None,
        actor_label=actor_label or (actor.full_name if actor else "system"),
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        details=details,
    )
    session.add(row)
    await session.flush()
    return row


async def seed_default_team(session: AsyncSession) -> Team:
    team = await session.scalar(select(Team).where(Team.name == "АМПасадори"))
    if not team:
        try:
            async with session.begin_nested():
                session.add(Team(name="АМПасадори", description="Основна волонтерська команда Анисівського молодіжного простору"))
                await session.flush()
        except IntegrityError:
            pass
        team = await session.scalar(select(Team).where(Team.name == "АМПасадори"))
    if not team:
        raise RuntimeError("Не вдалося створити або отримати команду АМПасадори")
    return team


async def add_active_users_to_default_team(session: AsyncSession) -> None:
    team = await seed_default_team(session)
    users = (await session.scalars(select(User).where(User.status == UserStatus.ACTIVE.value))).all()
    for user in users:
        exists = await session.scalar(select(TeamMember).where(TeamMember.team_id == team.id, TeamMember.user_id == user.id))
        if exists:
            continue
        try:
            async with session.begin_nested():
                session.add(TeamMember(team_id=team.id, user_id=user.id))
                await session.flush()
        except IntegrityError:
            # Another worker attached the same user to the default team.
            pass


async def complete_team_quest(session: AsyncSession, quest: Quest, admin_user: User | None = None) -> int:
    if quest.quest_type != "team" or quest.completed:
        return 0
    if quest.progress_value < quest.target_value:
        return 0

    # v1.5: a team quest has an explicit participant list. Reward only people
    # who actually joined the quest. For legacy team quests without joins, fall
    # back to the default team so older data keeps working.
    parts = (await session.scalars(
        select(QuestParticipation).where(
            QuestParticipation.quest_id == quest.id,
            QuestParticipation.status.in_(["joined", "completed"]),
        )
    )).all()
    user_ids = [p.user_id for p in parts]
    if not user_ids and quest.team_id:
        members = (await session.scalars(select(TeamMember).where(TeamMember.team_id == quest.team_id))).all()
        user_ids = [m.user_id for m in members]

    quest.xp_reward = normalize_quest_xp(quest.xp_reward, "team")
    count = 0
    for user_id in dict.fromkeys(user_ids):
        user = await session.get(User, user_id)
        if not user or user.status != UserStatus.ACTIVE.value:
            continue
        await add_xp(
            session, user, quest.xp_reward, f"Командний квест «{quest.title}»",
            category="team_quest", created_by=admin_user.id if admin_user else None,
        )
        part = next((p for p in parts if p.user_id == user_id), None)
        if part:
            part.status = "approved"
            if not part.completed_at:
                part.completed_at = datetime.utcnow()
            part.approved_at = datetime.utcnow()
        count += 1
    quest.completed = True
    quest.active = False
    await log_audit(session, "team_quest_completed", admin_user, entity_type="quest", entity_id=quest.id, details=f"Нагороджено: {count}")
    return count


def build_profile_qr_png(
    user: User,
    bot_username: str,
    *,
    total_xp: int = 0,
    level: str = "АМПасадор",
    profile_photo: bytes | None = None,
) -> bytes:
    """Create the print-ready personal QR badge (55 × 85 mm at 300 DPI)."""
    token = user.public_token or "missing"
    link = f"https://t.me/{bot_username}?start=profile_{token}"

    # 55 × 85 mm at 300 DPI = ~650 × 1004 px.
    W, H = 650, 1004
    canvas = Image.new("RGB", (W, H), "#EEF9FA")
    draw = ImageDraw.Draw(canvas)

    try:
        f_brand = ImageFont.truetype("DejaVuSans-Bold.ttf", 34)
        f_name = ImageFont.truetype("DejaVuSans-Bold.ttf", 31)
        f_id = ImageFont.truetype("DejaVuSans-Bold.ttf", 22)
        f_mid = ImageFont.truetype("DejaVuSans-Bold.ttf", 19)
        f_text = ImageFont.truetype("DejaVuSans.ttf", 17)
        f_small = ImageFont.truetype("DejaVuSans.ttf", 14)
        f_tiny = ImageFont.truetype("DejaVuSans.ttf", 12)
    except OSError:
        f_brand = f_name = f_id = f_mid = f_text = f_small = f_tiny = ImageFont.load_default()

    def fit_font(text: str, max_width: int, start_size: int, min_size: int = 10, *, bold: bool = True):
        """Return the largest DejaVu font that keeps text inside max_width."""
        family = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
        for size in range(start_size, min_size - 1, -1):
            try:
                font = ImageFont.truetype(family, size)
            except OSError:
                return f_small
            if draw.textbbox((0, 0), text, font=font)[2] <= max_width:
                return font
        try:
            return ImageFont.truetype(family, min_size)
        except OSError:
            return f_small

    # Outer badge and top brand panel.
    draw.rounded_rectangle((12, 12, W - 12, H - 12), radius=34, fill="white", outline="#BFE5E9", width=3)
    draw.rounded_rectangle((12, 12, W - 12, 290), radius=34, fill="#0B5B6C")
    draw.rectangle((12, 235, W - 12, 290), fill="#0B5B6C")
    draw.ellipse((480, -40, 720, 200), fill="#0AA8B6")
    draw.ellipse((-75, 835, 165, 1075), fill="#DDF5F7")

    # Brand logo on white plate.
    logo_path = Path("app/web/static/amp_logo.png")
    if logo_path.exists():
        logo = Image.open(logo_path).convert("RGBA")
        logo.thumbnail((240, 100), Image.Resampling.LANCZOS)
        plate_x = (W - logo.width - 34) // 2
        draw.rounded_rectangle((plate_x, 36, plate_x + logo.width + 34, 36 + logo.height + 22), radius=18, fill="white")
        canvas.paste(logo, (plate_x + 17, 47), logo)
    else:
        text = "АМПасадори"
        tw = draw.textbbox((0, 0), text, font=f_brand)[2]
        draw.text(((W - tw) / 2, 58), text, font=f_brand, fill="white")

    # Participant photo. The user photo is optional and is never cropped destructively outside the circular avatar.
    avatar_box = (52, 172, 216, 336)
    if profile_photo:
        try:
            photo = Image.open(BytesIO(profile_photo))
            photo = ImageOps.exif_transpose(photo).convert("RGB")
            photo = ImageOps.fit(photo, (156, 156), method=Image.Resampling.LANCZOS, centering=(0.5, 0.45))
            mask = Image.new("L", (156, 156), 0)
            ImageDraw.Draw(mask).ellipse((0, 0, 155, 155), fill=255)
            draw.ellipse((48, 168, 220, 340), fill="white")
            canvas.paste(photo, (56, 176), mask)
        except Exception:
            profile_photo = None
    if not profile_photo:
        draw.ellipse(avatar_box, fill="#D9F2F4", outline="white", width=5)
        initials = "".join(part[:1].upper() for part in (user.full_name or "АМП").split()[:2]) or "АМП"
        tw = draw.textbbox((0, 0), initials, font=f_brand)[2]
        th = draw.textbbox((0, 0), initials, font=f_brand)[3]
        draw.text(((avatar_box[0] + avatar_box[2] - tw) / 2, (avatar_box[1] + avatar_box[3] - th) / 2 - 4), initials, font=f_brand, fill="#0B5B6C")

    # Name and identity block.
    name = (user.full_name or "Учасник АМП").strip()
    max_name_w = 355
    words = name.split()
    lines, current = [], ""
    for word in words:
        candidate = (current + " " + word).strip()
        if not current or draw.textbbox((0, 0), candidate, font=f_name)[2] <= max_name_w:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    lines = lines[:2]
    y = 178
    for line in lines:
        draw.text((245, y), line, font=f_name, fill="white")
        y += 39
    draw.rounded_rectangle((245, 264, 405, 306), radius=18, fill="#E7F8F9")
    draw.text((262, 274), f"АМП-{user.id:04d}", font=f_id, fill="#0B5B6C")
    level_clean = "".join(ch for ch in level if ch.isalpha() or ch.isspace() or ch in "—-").strip() or "АМПасадор"
    level_top_font = fit_font(level_clean, 190, 14, 9, bold=False)
    level_top_w = draw.textbbox((0, 0), level_clean, font=level_top_font)[2]
    draw.text((420 + (190 - level_top_w) / 2, 276), level_clean, font=level_top_font, fill="white")

    # QR panel.
    draw.rounded_rectangle((50, 372, 600, 776), radius=30, fill="#F8FDFE", outline="#BFE5E9", width=2)
    draw.text((0, 0), "", font=f_small, fill="#0B5B6C")
    title = "ПЕРСОНАЛЬНИЙ QR-БЕЙДЖ"
    tw = draw.textbbox((0, 0), title, font=f_mid)[2]
    draw.text(((W - tw) / 2, 396), title, font=f_mid, fill="#0B5B6C")

    qr = qrcode.QRCode(version=None, box_size=10, border=2, error_correction=qrcode.constants.ERROR_CORRECT_M)
    qr.add_data(link)
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="#0B5B6C", back_color="white").convert("RGB").resize((292, 292), Image.Resampling.NEAREST)
    qr_x = (W - 292) // 2
    canvas.paste(qr_img, (qr_x, 438))
    hint = "Відскануй у Telegram — приватні дані не показуються"
    tw = draw.textbbox((0, 0), hint, font=f_tiny)[2]
    draw.text(((W - tw) / 2, 742), hint, font=f_tiny, fill="#607A80")

    # Compact metrics row.
    stat_y = 812
    stat_w, gap = 168, 14
    stat_x = 52
    stats = [
        ("XP", str(int(total_xp or 0))),
        ("ГОДИНИ", f"{float(user.volunteer_hours or 0):g}"),
        ("РІВЕНЬ", level_clean),
    ]
    for label_text, value in stats:
        draw.rounded_rectangle((stat_x, stat_y, stat_x + stat_w, stat_y + 90), radius=18, fill="#EAF8FA")
        lw = draw.textbbox((0, 0), label_text, font=f_tiny)[2]
        value_font = fit_font(value, stat_w - 22, 19, 10, bold=True)
        vw = draw.textbbox((0, 0), value, font=value_font)[2]
        draw.text((stat_x + (stat_w - lw) / 2, stat_y + 14), label_text, font=f_tiny, fill="#668188")
        draw.text((stat_x + (stat_w - vw) / 2, stat_y + 42), value, font=value_font, fill="#173B43")
        stat_x += stat_w + gap

    footer = "АМПасадори • Анисівський молодіжний простір"
    tw = draw.textbbox((0, 0), footer, font=f_tiny)[2]
    draw.text(((W - tw) / 2, 940), footer, font=f_tiny, fill="#0B5B6C")
    draw.rounded_rectangle((185, 970, 465, 976), radius=3, fill="#06B8C5")

    out = BytesIO()
    canvas.save(out, format="PNG", dpi=(300, 300), optimize=True)
    return out.getvalue()



def export_event_participants_pdf(
    event: Event,
    registrations: list[tuple[EventRegistration, User]],
    *,
    include_sensitive: bool = False,
) -> bytes:
    """Build a print-friendly branded PDF register for one event.

    The default document is deliberately data-minimized.  Sensitive contact,
    gender and vulnerability fields are added only for a superadmin route.
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from textwrap import shorten

    plt.rcParams["font.family"] = "DejaVu Sans"
    out = BytesIO()
    logo_path = Path("app/web/static/amp_logo.png")
    logo = plt.imread(str(logo_path)) if logo_path.exists() else None
    privacy = (
        "КОНФІДЕНЦІЙНО • розширений список • доступ лише суперадміністратору"
        if include_sensitive else
        "Внутрішній робочий список • без контактних і чутливих соціальних даних"
    )
    if include_sensitive:
        headers = ["№", "ПІБ", "Вік", "Стать", "Соціальний статус", "Email", "Телефон", "Фото/відео", "Статус", "ID АМП"]
        widths = [0.035, 0.16, 0.045, 0.075, 0.18, 0.14, 0.10, 0.09, 0.095, 0.08]
    else:
        headers = ["№", "ПІБ", "Вік", "Статус участі", "ID АМП"]
        widths = [0.06, 0.40, 0.09, 0.27, 0.18]

    rows_per_page = 22 if include_sensitive else 28
    chunks = [registrations[i:i + rows_per_page] for i in range(0, len(registrations), rows_per_page)] or [[]]
    with PdfPages(out) as pdf:
        for page_no, chunk in enumerate(chunks, start=1):
            fig = plt.figure(figsize=(11.69, 8.27), facecolor="white")
            # branded masthead
            fig.patches.extend([plt.Rectangle((0, .89), 1, .11, transform=fig.transFigure, color="#0B5B6C", zorder=-1)])
            if logo is not None:
                axl = fig.add_axes([.035, .905, .075, .07]); axl.imshow(logo); axl.axis("off")
            fig.text(.125, .948, "АМПасадори", fontsize=20, weight="bold", color="white", va="center")
            fig.text(.125, .916, "Список учасників події", fontsize=11.5, color="#DDF8FA", va="center")
            fig.text(.965, .948, f"{page_no}/{len(chunks)}", fontsize=9, color="white", ha="right", va="center")

            fig.text(.04, .845, event.title, fontsize=16, weight="bold", color="#173B43")
            meta = f"{event.starts_at.strftime('%d.%m.%Y • %H:%M')}   •   {event.location or 'Локацію не зазначено'}   •   {len(registrations)} реєстрацій"
            fig.text(.04, .812, meta, fontsize=9.5, color="#607C83")
            fig.text(.04, .782, privacy, fontsize=8.2, color="#7A9095")

            ax = fig.add_axes([.035, .09, .93, .66]); ax.axis("off")
            body=[]
            for i, (reg, user) in enumerate(chunk, start=(page_no-1)*rows_per_page+1):
                age = age_on(user.birth_date, event.starts_at.date()) if user.birth_date else "—"
                if include_sensitive:
                    vulnerabilities = ", ".join(vulnerability_labels(user.vulnerability_categories)) or "—"
                    body.append([
                        i,
                        shorten(user.full_name or "—", width=32, placeholder="…"),
                        age,
                        gender_label(user.gender) if user.gender else "—",
                        shorten(vulnerabilities, width=45, placeholder="…"),
                        shorten(user.email or "—", width=28, placeholder="…"),
                        user.phone or "—",
                        media_consent_label(user.media_consent),
                        event_registration_status_label(reg.status),
                        f"АМП-{user.id:04d}",
                    ])
                else:
                    body.append([i, user.full_name or "—", age, event_registration_status_label(reg.status), f"АМП-{user.id:04d}"])
            if body:
                table = ax.table(cellText=body, colLabels=headers, cellLoc="left", colLoc="left", loc="upper left", colWidths=widths)
                table.auto_set_font_size(False); table.set_fontsize(7.2 if include_sensitive else 8.6); table.scale(1, 1.42)
                for (r,c), cell in table.get_celld().items():
                    cell.set_edgecolor("#D4E6E9")
                    cell.set_linewidth(.55)
                    cell.PAD=.12
                    if r == 0:
                        cell.set_facecolor("#EAF8FA"); cell.set_text_props(weight="bold", color="#0B5B6C")
                    elif r % 2 == 0:
                        cell.set_facecolor("#F7FBFC")
                    else:
                        cell.set_facecolor("white")
            else:
                ax.text(.5,.55,"На подію ще ніхто не зареєструвався",ha="center",va="center",fontsize=13,color="#71888E")
            fig.text(.04,.035,"АМП • Анисівський молодіжний простір",fontsize=8,color="#0B5B6C")
            fig.text(.96,.035,f"Сформовано {datetime.now().strftime('%d.%m.%Y %H:%M')}",fontsize=7.5,color="#71888E",ha="right")
            pdf.savefig(fig, bbox_inches="tight", pad_inches=.03); plt.close(fig)
    return out.getvalue()


def export_event_participants_excel(
    event: Event,
    registrations: list[tuple[EventRegistration, User]],
    *,
    include_sensitive: bool = False,
) -> bytes:
    """Build an event-specific participant register.

    The default workbook is data-minimized and contains only operational fields.
    Sensitive contact, gender and vulnerability data are included only when
    ``include_sensitive=True`` and the web route has already enforced
    superadmin access.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Учасники події"
    ws.sheet_view.showGridLines = False

    dark_teal = "0B5B6C"
    turquoise = "06B8C5"
    light_teal = "EAF8FA"
    pale = "F5FBFC"
    white = "FFFFFF"
    text = "173B43"
    muted = "6C858B"
    line = "CFE5E9"

    thin = Side(style="thin", color=line)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    if include_sensitive:
        headers = [
            "№", "ПІБ", "Вік", "Стать", "Соціальний статус",
            "Електронна пошта", "Телефон", "Згода на фото/відеозйомку",
            "Статус участі", "ID АМП", "Підтверджено", "Цифровий код підтвердження",
        ]
        note = "Конфіденційно • розширений експорт • лише для суперадміністратора"
        widths = {"A": 6, "B": 26, "C": 8, "D": 18, "E": 36, "F": 28, "G": 18, "H": 25, "I": 22, "J": 13, "K": 19, "L": 68}
    else:
        headers = ["№", "ПІБ", "Вік", "Статус участі", "ID АМП", "Підтверджено", "Цифровий код підтвердження"]
        note = "Внутрішній робочий список • без чутливих контактних і соціальних даних"
        widths = {"A": 6, "B": 32, "C": 10, "D": 25, "E": 15, "F": 19, "G": 68}

    max_col = len(headers)
    end_col_letter = get_column_letter(max_col)

    ws.merge_cells(f"A1:{end_col_letter}1")
    ws["A1"] = "Список учасників події"
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=dark_teal)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(f"A2:{end_col_letter}2")
    ws["A2"] = note
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color=muted)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].fill = PatternFill("solid", fgColor=pale)

    meta = [
        ("Назва події", event.title),
        ("Дата", event.starts_at.strftime("%d.%m.%Y")),
        ("Час", event.starts_at.strftime("%H:%M")),
        ("Місце", event.location or "—"),
        ("Кількість реєстрацій", len(registrations)),
    ]
    for idx, (label_text, value) in enumerate(meta, start=4):
        ws[f"A{idx}"] = label_text
        ws[f"A{idx}"].font = Font(name="Arial", size=11, bold=True, color=dark_teal)
        ws[f"A{idx}"].fill = PatternFill("solid", fgColor=light_teal)
        ws[f"A{idx}"].border = border
        ws[f"A{idx}"].alignment = Alignment(vertical="center")
        if max_col > 1:
            ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=max_col)
            cell = ws.cell(idx, 2)
            cell.value = value
            cell.font = Font(name="Arial", size=11, color=text)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            for col in range(2, max_col + 1):
                ws.cell(idx, col).border = border
                ws.cell(idx, col).fill = PatternFill("solid", fgColor=white)
        ws.row_dimensions[idx].height = 24

    header_row = 10
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, value)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=turquoise)
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 34

    status_labels = {
        "registered": "Зареєстрований",
        "reserved": "Місце зарезервовано",
        "waitlisted": "У черзі",
        "checked_in": "Відмічено присутність",
        "attended": "Був присутній",
        "no_show": "Не прийшов",
        "cancelled": "Скасував",
    }
    event_day = event.starts_at.date()
    for number, (reg, user) in enumerate(registrations, start=1):
        row = header_row + number
        if include_sensitive:
            values = [
                number,
                user.full_name or "—",
                age_on(user.birth_date, event_day) if user.birth_date else None,
                gender_label(user.gender),
                "; ".join(vulnerability_labels(user.vulnerability_categories)) or "Не зазначено",
                user.email or "—",
                user.phone or "—",
                media_consent_label(user.media_consent),
                status_labels.get(reg.status, reg.status or "—"),
                f"АМП-{user.id:04d}",
                reg.confirmed_at.strftime("%d.%m.%Y %H:%M") if reg.confirmed_at else "—",
                reg.attendance_signature or "—",
            ]
        else:
            values = [
                number,
                user.full_name or "—",
                age_on(user.birth_date, event_day) if user.birth_date else None,
                status_labels.get(reg.status, reg.status or "—"),
                f"АМП-{user.id:04d}",
                reg.confirmed_at.strftime("%d.%m.%Y %H:%M") if reg.confirmed_at else "—",
                reg.attendance_signature or "—",
            ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.font = Font(name="Arial", size=10, color=text)
            cell.border = border
            cell.fill = PatternFill("solid", fgColor=white if number % 2 else pale)
            cell.alignment = Alignment(
                horizontal="center" if col in ({1, 3, 4, 8, 9, 10, 11} if include_sensitive else {1, 3, 4, 5, 6}) else "left",
                vertical="center",
                wrap_text=True,
            )
        ws.row_dimensions[row].height = 30

    last_row = max(header_row, header_row + len(registrations))
    ws.auto_filter.ref = f"A{header_row}:{end_col_letter}{last_row}"
    ws.freeze_panes = f"A{header_row + 1}"
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = turquoise
    ws.page_margins.left = 0.35; ws.page_margins.right = 0.35; ws.page_margins.top = 0.55; ws.page_margins.bottom = 0.55
    ws.oddHeader.center.text = "&BАМПасадори • Список учасників події"
    ws.oddHeader.center.size = 10
    ws.oddFooter.left.text = "АМП • Анисівський молодіжний простір"
    ws.oddFooter.right.text = "Сторінка &P з &N"
    ws.sheet_view.zoomScale = 90
    wb.properties.title = f"Учасники події — {event.title}"
    wb.properties.subject = "АМПасадори • реєстр учасників події"
    wb.properties.creator = "АМП XP / АМПасадори"

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


async def export_basic_excel(session: AsyncSession) -> bytes:
    """Data-minimized operational export available to regular web admins."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Учасники"
    ws.append(["ID АМП", "ПІБ", "Вік", "Статус", "Роль", "Населений пункт", "XP загальний", "XP сезону", "Волонтерські години"])
    users = (await session.scalars(select(User).order_by(User.full_name.asc()))).all()
    season = await current_season(session)
    for u in users:
        ws.append([
            f"АМП-{u.id:04d}",
            u.full_name,
            age_on(u.birth_date) if u.birth_date else None,
            u.status,
            u.role,
            u.settlement,
            await xp_total(session, u.id),
            await season_xp(session, u.id, season.id if season else None),
            u.volunteer_hours,
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(1, ws.max_row)}"
    for col, width in {"A": 14, "B": 32, "C": 8, "D": 16, "E": 18, "F": 24, "G": 14, "H": 14, "I": 20}.items():
        ws.column_dimensions[col].width = width

    ws2 = wb.create_sheet("Події")
    ws2.append(["ID", "Назва", "Дата", "Місце", "XP", "Години", "Статус"])
    events = (await session.scalars(select(Event).order_by(Event.starts_at.desc()))).all()
    for e in events:
        ws2.append([e.id, e.title, e.starts_at, e.location, e.xp_reward, e.volunteer_hours, e.status])

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


async def export_excel(session: AsyncSession) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Учасники"
    ws.append(["ID", "ID Telegram", "Прізвище", "Ім’я", "ПІБ", "Дата народження", "Вік", "Стать", "Телефон", "Email", "Telegram", "Статус", "Роль", "Населений пункт", "Соціальний статус / категорії вразливості", "Згода на фото/відео", "XP загальний", "XP сезону", "XP-гаманець", "Волонтерські години", "Код запрошення", "Ознайомлення з даними — версія", "Ознайомлення — дата", "Блокування до", "Причина блокування"])
    users = (await session.scalars(select(User).order_by(User.id))).all()
    season = await current_season(session)
    for u in users:
        first_name, last_name = split_display_name(u.full_name, u.first_name, u.last_name)
        ws.append([
            u.id,
            u.tg_id,
            last_name,
            first_name,
            u.full_name,
            u.birth_date,
            age_on(u.birth_date) if u.birth_date else None,
            gender_label(u.gender),
            u.phone,
            u.email,
            f"@{u.username}" if u.username else "",
            u.status,
            u.role,
            u.settlement,
            "; ".join(vulnerability_labels(u.vulnerability_categories)),
            media_consent_label(u.media_consent),
            await xp_total(session, u.id),
            await season_xp(session, u.id, season.id if season else None),
            u.wallet_xp,
            u.volunteer_hours,
            u.referral_code,
            u.privacy_notice_version,
            u.privacy_acknowledged_at,
            u.blocked_until,
            u.block_reason,
        ])

    ws2 = wb.create_sheet("XP журнал")
    ws2.append(["ID", "ID учасника", "XP", "Категорія", "Опис", "ID сезону", "Дата"])
    txs = (await session.scalars(select(XPTransaction).order_by(XPTransaction.created_at.desc()))).all()
    for t in txs:
        ws2.append([t.id, t.user_id, t.amount, t.category, t.description, t.season_id, t.created_at])

    ws3 = wb.create_sheet("Події")
    ws3.append(["ID", "Назва", "Дата", "Локація", "XP", "Години", "Статус"])
    events = (await session.scalars(select(Event).order_by(Event.starts_at.desc()))).all()
    for e in events:
        ws3.append([e.id, e.title, e.starts_at, e.location, e.xp_reward, e.volunteer_hours, e.status])

    ws4 = wb.create_sheet("Реферали")
    ws4.append(["ID", "ID запрошувача", "ID запрошеного", "Статус", "XP", "Дата"])
    refs = (await session.scalars(select(Referral).order_by(Referral.created_at.desc()))).all()
    for r in refs:
        ws4.append([r.id, r.inviter_user_id, r.invited_user_id, r.status, r.xp_reward, r.created_at])

    ws5 = wb.create_sheet("Бейджі")
    ws5.append(["ID учасника", "ID бейджа", "Дата"])
    awarded = (await session.scalars(select(UserBadge).order_by(UserBadge.awarded_at.desc()))).all()
    for row in awarded:
        ws5.append([row.user_id, row.badge_id, row.awarded_at])

    ws6 = wb.create_sheet("Винагороди")
    ws6.append(["ID", "Назва", "Вартість XP", "Залишок", "Активна", "Фото"])
    rewards = (await session.scalars(select(Reward).order_by(Reward.id))).all()
    for r in rewards:
        ws6.append([r.id, r.title, r.min_xp, r.stock, r.active, r.image_path])

    ws7 = wb.create_sheet("Заявки на винагороди")
    ws7.append(["ID", "ID винагороди", "ID учасника", "Статус", "XP витрачено", "Дата заявки", "Дата видачі"])
    claims = (await session.scalars(select(RewardClaim).order_by(RewardClaim.requested_at.desc()))).all()
    for c in claims:
        ws7.append([c.id, c.reward_id, c.user_id, c.status, c.xp_spent, c.requested_at, c.fulfilled_at])

    ws8 = wb.create_sheet("Волонтерські задачі")
    ws8.append(["ID", "Назва", "XP", "Години", "Дедлайн", "Статус", "Макс. учасників", "Фото"])
    tasks = (await session.scalars(select(VolunteerTask).order_by(VolunteerTask.created_at.desc()))).all()
    for t in tasks:
        ws8.append([t.id, t.title, t.xp_reward, t.hours_reward, t.deadline, t.status, t.max_participants, t.image_path])

    ws8b = wb.create_sheet("Участь у задачах")
    ws8b.append(["ID", "ID задачі", "ID учасника", "Статус", "Долучився", "Подано", "Підтверджено", "Коментар"])
    task_parts = (await session.scalars(select(VolunteerTaskParticipation).order_by(VolunteerTaskParticipation.joined_at.desc()))).all()
    for p in task_parts:
        ws8b.append([p.id, p.task_id, p.user_id, p.status, p.joined_at, p.submitted_at, p.approved_at, p.admin_note])

    ws9 = wb.create_sheet("Ідеї")
    ws9.append(["ID", "ID учасника", "Назва", "Напрям", "Проблема", "Рішення", "Аудиторія", "Результат", "Ресурси", "Статус", "Відповідальний", "Примітка", "Створено", "Оновлено"])
    ideas = (await session.scalars(select(Idea).order_by(Idea.created_at.desc()))).all()
    for i in ideas:
        ws9.append([i.id, i.user_id, i.title, i.category, i.problem, i.description, i.audience, i.expected_result, i.resources, i.status, i.responsible_user_id, i.admin_note, i.created_at, i.updated_at])

    ws10 = wb.create_sheet("Звернення")
    ws10.append(["ID", "ID учасника", "Тип", "Тема", "Опис", "Пріоритет", "Статус", "Відповідальний", "Відповідь", "Внутрішня примітка", "Фото", "Створено", "Оновлено", "Вирішено"])
    cases = (await session.scalars(select(RequestCase).order_by(RequestCase.created_at.desc()))).all()
    for c in cases:
        ws10.append([c.id, c.user_id, c.category, c.title, c.description, c.priority, c.status, c.assigned_user_id, c.admin_response, c.internal_note, c.image_path, c.created_at, c.updated_at, c.resolved_at])

    ws11 = wb.create_sheet("Каталог активностей")
    ws11.append(["ID", "Код", "Назва", "Категорія", "XP", "Години", "Активна", "Пояснення"])
    activity_types = (await session.scalars(select(ActivityType).order_by(ActivityType.sort_order, ActivityType.id))).all()
    for a in activity_types:
        ws11.append([a.id, a.code, a.title, a.category, a.xp_reward, a.hours_reward, a.active, a.description])

    ws12 = wb.create_sheet("Заявки на активності")
    ws12.append(["ID", "ID активності", "ID учасника", "Статус", "План", "Результат", "XP", "Години", "Подано", "Схвалено", "Виконано"])
    apps = (await session.scalars(select(ActivityApplication).order_by(ActivityApplication.requested_at.desc()))).all()
    for a in apps:
        ws12.append([a.id, a.activity_type_id, a.user_id, a.status, a.plan_text, a.result_note, a.xp_reward, a.hours_reward, a.requested_at, a.approved_at, a.completed_at])

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
