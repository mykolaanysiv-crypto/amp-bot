"""Regression gates for document word wrapping and Ukrainian report labels."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from app.document_layout import apply_excel_word_wrap, enable_docx_word_wrap, word_lines

ROOT = Path(__file__).resolve().parents[1]


def test_wrap_breaks_words_not_letters_and_grows_excel_rows():
    sample = "Документ містить довгу назву молодіжної активності для нашої громади"
    lines = word_lines(sample, 25)
    assert len(lines) >= 3
    assert " ".join(lines) == sample
    wb = Workbook()
    sheet = wb.active
    sheet.column_dimensions['A'].width = 28
    sheet['A1'] = 'Назва заходу'
    sheet['A2'] = sample
    sheet['A2'].alignment = Alignment(horizontal='left')
    apply_excel_word_wrap(sheet)
    assert sheet['A2'].alignment.wrap_text is True
    assert sheet.row_dimensions[2].height > sheet.row_dimensions[1].height
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    restored = load_workbook(bio)
    assert restored.active['A2'].alignment.wrap_text is True
    assert restored.active.row_dimensions[2].height > 35


def test_word_table_removes_fixed_row_height_and_no_wrap():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[1].height = Pt(14)
    table.cell(1, 0).text = 'Ініціатива для молоді з багатьма словами та детальним описом'
    enable_docx_word_wrap(doc)
    assert table.autofit is True
    assert table.rows[1].height is None
    out = BytesIO()
    doc.save(out)
    assert out.getvalue().startswith(b'PK')


def test_quest_field_labels_and_ukrainian_analytics():
    html = (ROOT / 'app/web/templates/quests.html').read_text()
    for label in ('Спільна ціль команди', 'Подія для QR-квесту', 'Допустиме запізнення'):
        assert label in html
    styles = (ROOT / 'app/web/static/admin.css').read_text()
    assert '.quest-auto-field[hidden]' in styles
    for rel in ('app/web/templates/dashboard.html', 'app/web/templates/event_detail.html',
                'app/reporting/exports.py', 'app/web/templates/help.html', 'app/analytics_modules/core.py'):
        assert 'воронк' not in (ROOT / rel).read_text().lower(), rel
