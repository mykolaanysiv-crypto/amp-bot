from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO

from docx import Document
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.event_documents import fill_registration_template
from app.models import EventRegistration, Referral, Reward, UserRole, UserStatus
from app.services import (
    add_xp,
    create_event,
    register_for_event,
    revoke_referral_reward_if_inactive,
    seed_default_space_rewards,
    xp_total,
)
from tests.conftest import create_user


async def test_referral_reward_clawback_within_30_days_is_idempotent(db):
    async with db.session_factory() as session:
        inviter = await create_user(session, tg_id=174101, name="Запрошувач Тест")
        invited = await create_user(session, tg_id=174102, name="Запрошений Тест")
        referral = Referral(
            inviter_user_id=inviter.id,
            invited_user_id=invited.id,
            status="rewarded",
            xp_reward=8,
            rewarded_at=datetime.utcnow() - timedelta(days=10),
        )
        session.add(referral)
        await add_xp(session, inviter, 8, "Тестовий реферальний бонус", category="referral")
        await session.flush()
        assert await xp_total(session, inviter.id) == 8
        assert inviter.wallet_xp == 8

        invited.status = UserStatus.INACTIVE.value
        result = await revoke_referral_reward_if_inactive(session, invited, now=datetime.utcnow())
        assert result is not None
        returned_inviter, removed_xp, days_after = result
        assert returned_inviter.id == inviter.id
        assert removed_xp == 8
        assert 9 <= days_after <= 10
        assert referral.status == "revoked"
        assert referral.clawback_xp == 8
        assert await xp_total(session, inviter.id) == 0
        assert inviter.wallet_xp == 0

        # A repeated inactive/status callback must not remove XP twice.
        assert await revoke_referral_reward_if_inactive(session, invited, now=datetime.utcnow()) is None
        assert await xp_total(session, inviter.id) == 0


async def test_default_space_rewards_seeded_as_repeatable_services(db):
    async with db.session_factory() as session:
        await seed_default_space_rewards(session)
        await seed_default_space_rewards(session)
        await session.flush()
        rows = list((await session.scalars(select(Reward).where(Reward.reward_type == "service").order_by(Reward.min_xp))).all())
        assert len(rows) == 9
        assert [r.min_xp for r in rows] == [5, 15, 45, 50, 50, 75, 100, 100, 150]
        assert all(r.stock is None and r.active for r in rows)


async def test_donor_xlsx_template_preserves_layout_and_fills_confirmation_code(db):
    async with db.session_factory() as session:
        admin = await create_user(session, tg_id=174111, name="Адмін Тест", role=UserRole.ADMIN.value)
        user = await create_user(session, tg_id=174112, name="Учасник Тест")
        user.last_name = "Тест"
        user.first_name = "Учасник"
        event = await create_event(session, "Подія донора", "", datetime.utcnow() + timedelta(days=1), "АМП", 10, 0, admin.id)
        reg = await register_for_event(session, user.id, event.id)
        reg.status = "attended"
        reg.confirmed_at = datetime.utcnow()
        reg.attendance_signature = "a" * 64
        await session.flush()

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "ФОРМА ДОНОРА"
        ws["A2"] = "Подія: {{event_title}}"
        ws.append(["№", "ПІБ", "Статус участі", "Цифровий код підтвердження"])
        ws.append(["", "", "", ""])
        ws["A4"].font = ws["A1"].font.copy(bold=True)
        original = BytesIO(); wb.save(original)

        data, _, ext = fill_registration_template(original.getvalue(), "donor.xlsx", event, [(reg, user)], include_sensitive=False)
        assert ext == ".xlsx"
        out = load_workbook(BytesIO(data))
        out_ws = out.active
        assert out_ws["A1"].value == "ФОРМА ДОНОРА"
        assert out_ws["A2"].value == f"Подія: {event.title}"
        assert out_ws["A4"].value == 1
        assert out_ws["B4"].value == user.full_name
        assert out_ws["D4"].value == "a" * 64


async def test_donor_docx_template_fills_existing_table_and_keeps_heading(db):
    async with db.session_factory() as session:
        admin = await create_user(session, tg_id=174121, name="Адмін Word", role=UserRole.ADMIN.value)
        user = await create_user(session, tg_id=174122, name="Учасник Word")
        event = await create_event(session, "Word-подія", "", datetime.utcnow() + timedelta(days=2), "АМП", 10, 0, admin.id)
        reg = await register_for_event(session, user.id, event.id)
        reg.status = "attended"
        reg.confirmed_at = datetime.utcnow()
        reg.attendance_signature = "b" * 64
        await session.flush()

        doc = Document()
        doc.add_heading("РЕЄСТРАЦІЙНИЙ ЛИСТ ДОНОРА", level=1)
        doc.add_paragraph("Назва: {{event_title}}")
        table = doc.add_table(rows=2, cols=4)
        for idx, text in enumerate(["№", "ПІБ", "Статус", "Підпис"]):
            table.rows[0].cells[idx].text = text
        buf = BytesIO(); doc.save(buf)

        data, _, ext = fill_registration_template(buf.getvalue(), "donor.docx", event, [(reg, user)], include_sensitive=False)
        assert ext == ".docx"
        out = Document(BytesIO(data))
        assert out.paragraphs[0].text == "РЕЄСТРАЦІЙНИЙ ЛИСТ ДОНОРА"
        assert out.paragraphs[1].text == f"Назва: {event.title}"
        values = [cell.text for cell in out.tables[0].rows[1].cells]
        assert values[0] == "1"
        assert values[1] == user.full_name
        assert values[3] == "b" * 64
